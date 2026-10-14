"""Focused integration, expected2s/max20s: source review is real, attributed and exact-bound."""
import copy,hashlib,json
import pytest
from catalog_json.manual_review import source_review_update
from catalog_json.schema import PageContent,TextBlock,Figure,PageContext,ContextEvidence
from catalog_json.tools import write_json


def case(tmp_path):
    content=PageContent(title='Actual catalogue',text_blocks=[TextBlock(kind='paragraph',text='Read original')],tables=[],figures=[])
    p=tmp_path/'candidate-1.json';write_json(p,content.model_dump());sha=hashlib.sha256(p.read_bytes()).hexdigest()
    origin={'status':'receipt_verified','source_sha256':'a'*64,'physical_page':1,'attempt':1,'artifact_sha256':sha,'model':'anthropic.MiniMax-M3.1-Flash-Preview','program':{'package_version':'1.5.17'},'contract_sha256':'b'*64}
    record={'source_sha256':'a'*64,'pdf_page':1,'candidate_sha256':sha,'candidate_producer':origin,'content':content.model_dump(),'status':'awaiting_review'}
    proof={'schema_version':'catalogue-source-review/1','status':'source_verified','source_sha256':'a'*64,'physical_page':1,'unchanged_native_candidate':True,'corrections':[],'unresolved_issues':[],'full_original_image_review':True,'independent_from_initial_extractor':True,'reviewer':{'harness':'codex','session_id':'actual-root-session','version':'root-r1'},'checks':{k:True for k in ['all_text_and_meaning','all_table_cells','headers_units_conditions','merged_cells_and_blanks','figures_and_relationships','source_provenance']},'basis_candidate':{'sha256':sha,'producer':copy.deepcopy(origin)},'tables':{'row_counts':[],'column_counts':[],'cells_checked':0},'source_findings':'Whole original personally inspected before candidate comparison; source text matches.'}
    write_json(tmp_path/'root-source-review.json',proof);return record,proof


def test_unchanged_native_candidate_gets_separately_attributed_codex_review(tmp_path):
    record,proof=case(tmp_path);before=copy.deepcopy(record);update=source_review_update(tmp_path,record)
    assert update['status']=='accepted' and update['review']['verdict']=='pass'
    assert update['review_producer']['status']=='manual_source_review_verified'
    assert 'Codex' in update['reviewer'] and not update['manual_source_review']['native_reviewer_call']
    assert record==before and update['review_producer']['reviewed_candidate_sha256']==record['candidate_sha256']


@pytest.mark.parametrize('change',['wrong_source','wrong_page','changed_bytes','new_attempt','audit_changed','veto','edited_candidate','missing_check','wrong_cells','unresolved','drawing','bad_json'])
def test_stale_or_ineligible_manual_review_does_not_close_native_work(tmp_path,change):
    record,proof=case(tmp_path)
    if change=='wrong_source':proof['source_sha256']='c'*64
    if change=='wrong_page':proof['physical_page']=2
    if change=='changed_bytes':(tmp_path/'candidate-1.json').write_text('{}')
    if change=='new_attempt':write_json(tmp_path/'candidate-2.json',record['content'])
    if change=='audit_changed':write_json(tmp_path/'external-audit.json',{'findings':[]})
    if change=='veto':
        write_json(tmp_path/'external-audit.json',{'rejected_candidate_sha256':[record['candidate_sha256']]});proof['external_audit_sha256']=hashlib.sha256((tmp_path/'external-audit.json').read_bytes()).hexdigest()
    if change=='edited_candidate':proof['unchanged_native_candidate']=False;proof['corrections']=['Root edited this; independent model review still needed']
    if change=='missing_check':proof['checks'].pop('merged_cells_and_blanks')
    if change=='wrong_cells':proof['tables']['cells_checked']=1
    if change=='unresolved':record['content']['uncertainties']=['Missing real source fact']
    if change=='drawing':
        record['content']['figures']=[Figure(title='Engineering drawing',description='D diameter',labels=['D']).model_dump()];write_json(tmp_path/'candidate-1.json',record['content']);sha=hashlib.sha256((tmp_path/'candidate-1.json').read_bytes()).hexdigest();record['candidate_sha256']=sha;record['candidate_producer']['artifact_sha256']=sha;proof['basis_candidate']['sha256']=sha;proof['basis_candidate']['producer']['artifact_sha256']=sha
    write_json(tmp_path/'root-source-review.json',proof)
    if change=='bad_json':(tmp_path/'root-source-review.json').write_text('{')
    assert source_review_update(tmp_path,record) is None


def test_export_keeps_manual_reviewer_and_missing_native_review_honest(tmp_path):
    from catalog_json.export import to_xlsx
    from openpyxl import load_workbook
    record,proof=case(tmp_path);update=source_review_update(tmp_path,record);record.update(update)
    to_xlsx({'complete_for_selection':True,'pages':[record]},tmp_path/'reviewed.xlsx')
    w=load_workbook(tmp_path/'reviewed.xlsx',read_only=True);row=list(w['Происхождение'].values)[1]
    assert 'Codex' in row[8] and row[9]=='root-r1';w.close()
    record['status']='awaiting_review';record['review_producer']=None
    to_xlsx({'complete_for_selection':False,'pages':[record]},tmp_path/'pending.xlsx',allow_unverified=True)
    w=load_workbook(tmp_path/'pending.xlsx',read_only=True);assert list(w['Происхождение'].values)[1][8]=='не подтверждён';w.close()
