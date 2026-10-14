"""Focused integration: real MCP transport/images and artifact validation; expected 4 s, max 20 s."""
import asyncio
import json
from pathlib import Path
import sys
import pymupdf
from mcp import ClientSession,StdioServerParameters
from mcp.client.stdio import stdio_client
from catalog_json.tools import CatalogTools
from test_contract import content,review


def pdf(tmp_path):
    source=tmp_path/'source.pdf'
    with pymupdf.open() as d:
        p=d.new_page();p.insert_text((70,70),'Catalog example');d.save(source)
    return source


def test_save_requires_inspection_and_returns_errors(tmp_path):
    t=CatalogTools(pdf(tmp_path),tmp_path/'output')
    assert not t.save(1,content().model_dump_json(),'page')['saved']
    t.render(1);t.direct(1)
    bad=t.save(1,'{"title":1}','page')
    assert not bad['saved'] and bad['validation_errors']
    assert t.save(1,content().model_dump_json(),'page')['saved']
    r=CatalogTools(t.source,t.output,role='reviewer');r.render(1)
    assert r.save(1,content().model_dump_json(),'observation')['saved'];r.read(1,'page')
    incomplete=review();incomplete.checks.pop()
    assert not r.save(1,incomplete.model_dump_json(),'review')['saved']
    assert not r.save(1,content().model_dump_json(),'page')['saved']
    assert r.save(1,review().model_dump_json(),'review')['saved']


def test_native_mcp_image_roundtrip(tmp_path):
    source=pdf(tmp_path)
    async def run():
        params=StdioServerParameters(command=sys.executable,args=['-m','catalog_json.tools','--source',str(source),'--output',str(tmp_path/'out'),'--model','native-test','--contract-sha256','c'*64])
        async with stdio_client(params) as (reader,writer):
            async with ClientSession(reader,writer) as client:
                await client.initialize();listed=await client.list_tools()
                assert 'save_page_json' in [t.name for t in listed.tools]
                skill=await client.call_tool('load_skill',{});assert 'Fidelity rules' in skill.content[0].text
                result=await client.call_tool('view_page',{'page':1})
                assert not result.is_error
                assert any(b.type=='image' for b in result.content)
                result=await client.call_tool('get_json_schema',{'kind':'page'})
                assert not result.is_error
                await client.call_tool('extract_direct',{'page':1})
                saved=await client.call_tool('save_page_json',{'page':1,'json_text':content().model_dump_json()});assert not saved.is_error
                from catalog_json.staged import validated_candidate
                work=tmp_path/'out/page-0001'
                import hashlib
                assert validated_candidate(work,hashlib.sha256(source.read_bytes()).hexdigest(),1,'native-test','c'*64)
                assert validated_candidate(work,hashlib.sha256(source.read_bytes()).hexdigest(),1,'native-test','d'*64) is None
    asyncio.run(asyncio.wait_for(run(),20))


def test_json_exports_keep_units_qualifiers_and_full_page(tmp_path):
    """Fast unit: exports must not drop header conditions or drawings; expected <1 s, max 5 s."""
    from catalog_json.schema import Figure,TextBlock
    from catalog_json.export import to_html,to_xlsx
    from openpyxl import load_workbook
    d=content();d.tables[0].columns[0].label='Масса';d.tables[0].columns[0].unit='кг';d.tables[0].columns[0].qualifier='не более'
    d.figures=[Figure(title='Вид А',description='Высота Н±2',labels=['Н±2'])];d.text_blocks=[TextBlock(kind='footer',text='Контакты')]
    catalog={'complete_for_selection':True,'pages':[{'pdf_page':1,'status':'accepted','content':d.model_dump()}]}
    to_html(catalog,tmp_path/'page.html');html=(tmp_path/'page.html').read_text();assert 'не более' in html and 'кг' in html and 'Н±2' in html
    to_xlsx(catalog,tmp_path/'page.xlsx');book=load_workbook(tmp_path/'page.xlsx');assert book['Стр 1 таб 1']['A1'].value=='Масса, кг; не более'
    assert any('Н±2' in str(row) for row in book['Текст и чертежи'].values)


def test_excel_shows_artifact_producers_not_new_queue_version(tmp_path):
    """Fast unit: public exports must not relabel historical data; <1s/max5s."""
    from catalog_json.export import to_xlsx
    from openpyxl import load_workbook
    page={'pdf_page':35,'status':'accepted','content':content().model_dump(),
          'candidate_producer':{'status':'receipt_verified','model':'MiniMax','program':{'package_version':'1.4.15','source_code_sha256':'extractor-sha'}},
          'review_producer':{'status':'receipt_verified','model':'GLM','program':{'package_version':'1.4.8','source_code_sha256':'reviewer-sha'}}}
    catalog={'complete_for_selection':True,'tested_program':{'package_version':'1.4.26'},'source':{'filename':'source.pdf','sha256':'source-sha'},'pages':[page]}
    to_xlsx(catalog,tmp_path/'export.xlsx');sheet=load_workbook(tmp_path/'export.xlsx')['Происхождение']
    row=dict(zip(next(sheet.values),list(sheet.values)[1]));assert row['PDF страница']==35
    assert row['Версия извлекателя']=='1.4.15' and row['Версия проверяющего']=='1.4.8'
    assert row['Исходный PDF']=='source.pdf' and row['Извлекатель']=='MiniMax'
    catalog['pages'][0].pop('candidate_producer');to_xlsx(catalog,tmp_path/'unknown.xlsx')
    unknown=load_workbook(tmp_path/'unknown.xlsx')['Происхождение'];row=dict(zip(next(unknown.values),list(unknown.values)[1]));assert row['Версия извлекателя']=='не подтверждена'


def test_bundled_skills_native_discovery():
    """Fast unit: prevent silently running without the user's native skill; expected <1 s, max 5 s."""
    from fast_agent.skills.registry import SkillRegistry
    from catalog_json.tools import SKILLS
    assert {m.name for m in SkillRegistry(directories=[SKILLS]).load_manifests()}=={'catalog-extract','catalog-review'}


def test_interrupted_review_resumes_without_reextracting(tmp_path):
    """Focused integration: don't repeat extraction after a review timeout; expected <1 s, max 5 s."""
    from catalog_json.pipeline import convert
    from catalog_json.tools import write_json
    class Workers:
        extractor=reviewer='test-model';contract_digest='test-contract';fail=True;calls=[]
        async def run(self,role,source,out,page,attempt,feedback=False):
            self.calls.append(role)
            if role=='reviewer':
                if self.fail:raise TimeoutError('review interrupted')
                return review()
            t=CatalogTools(source,out,page=page,attempt=attempt);t.render(page);t.direct(page)
            assert t.save(page,content().model_dump_json(),'page')['saved']
            return content()
    workers=Workers();source=pdf(tmp_path);out=tmp_path/'output'
    first=asyncio.run(convert(source,out,[1],workers));assert first['pages'][0]['status']=='failed'
    workers.fail=False
    final=asyncio.run(convert(source,out,[1],workers,resume=True))
    assert final['complete_for_selection'] and workers.calls==['extractor','reviewer','reviewer']


def test_false_pass_cannot_override_independent_values(tmp_path):
    """Fast unit: a confident judge must not certify mismatching cells; expected <1 s, max 5 s."""
    import pytest
    source=pdf(tmp_path);out=tmp_path/'out'
    t=CatalogTools(source,out);t.render(1);t.direct(1);bad=content();bad.tables[0].rows[0].cells[0].value=999
    assert t.save(1,bad.model_dump_json(),'page')['saved']
    r=CatalogTools(source,out,role='reviewer');r.render(1)
    with pytest.raises(ValueError):r.read(1,'page')
    assert r.save(1,content().model_dump_json(),'observation')['saved']
    disclosed=r.read(1,'page');assert disclosed['independent_comparison'][0]['source_observation']==80
    refused=r.save(1,review().model_dump_json(),'review');assert not refused['saved'] and refused['differences']
    assert not r.save(1,bad.model_dump_json(),'observation')['saved']


def test_independent_audit_catches_shift_and_missing_tolerance():
    """Fast unit: same dimensions/counts don't hide shifted values or diagram loss; expected <1 s, max 5 s."""
    from catalog_json.audit import compare_source
    from catalog_json.schema import Figure
    a=content();b=content();a.figures=[Figure(title='View',description='',labels=['H±2'])];b.figures=[Figure(title='View',description='H',labels=['H'])]
    assert any(d['path']=='figures[0].labels' for d in compare_source(a,b))


def test_minimax_preview_keeps_images_and_adaptive_reasoning():
    """Fast unit: prevent FastAgent silently stripping a supported model's image; expected <1 s, max 5 s."""
    from catalog_json.models import register_model_capabilities
    from fast_agent.llm.model_database import ModelDatabase
    register_model_capabilities()
    p=ModelDatabase.get_model_params('anthropic.MiniMax-M3.1-Flash-Preview?reasoning=low')
    assert p is not None and 'image/png' in p.tokenizes
    assert p.reasoning=='anthropic_thinking' and p.reasoning_effort_spec.allowed_efforts==['low','medium','high','xhigh','max']
    assert not p.anthropic_thinking_disable_supported


def test_observer_can_repair_own_error_only_after_new_source_view(tmp_path):
    """Fast unit: keep independence without freezing a mistaken observer forever; expected <1 s, max 5 s."""
    t=CatalogTools(pdf(tmp_path),tmp_path/'out');t.render(1);t.direct(1);assert t.save(1,content().model_dump_json(),'page')['saved']
    r=CatalogTools(t.source,t.output,role='reviewer');r.render(1);wrong=content();wrong.tables[0].rows[0].cells[0].value=999
    assert r.save(1,wrong.model_dump_json(),'observation')['saved'];r.read(1,'page')
    assert not r.correct_observation(1,content().model_dump_json(),'Checked source')['saved']
    r.render(1,[0,0,1,1]);assert r.correct_observation(1,content().model_dump_json(),'Fresh source image shows correct value')['saved']
    assert json.loads((r.work(1)/'source-observation-0.json').read_text())['tables'][0]['rows'][0]['cells'][0]['value']==999
    assert not r.read(1,'page')['independent_comparison']


def test_source_grid_expands_only_geometry_backed_merge(tmp_path):
    """Focused integration: distinguish empty from merged without guessing; expected <1 s, max 5 s."""
    from catalog_json.grid import table_grids
    with pymupdf.open() as d:
        p=d.new_page(width=300,height=300)
        p.draw_line((30,60),(270,60));p.draw_line((30,120),(150,120));p.draw_line((30,180),(270,180))
        p.insert_text((50,90),'row1');p.insert_text((50,150),'row2');p.insert_text((200,120),'100')
        result=table_grids(p,[.09,.19,.91,.61]);table=result['tables'][0]
        assert table['expanded_rows']==[['row1','100'],['row2','100']]
        assert table['merged_cells'][0]['covered_rows']==[0,1] and not table['unknown_cells']


def test_figure_grouping_preserves_labels_without_false_omission():
    """Fast unit: grouped views remain complete; reject lost numeric tolerances. Expected <1 s, max 5 s."""
    from catalog_json.audit import compare_source
    from catalog_json.schema import Figure
    a=content();b=content()
    a.figures=[Figure(title='Elevation',description='',labels=['H±2']),Figure(title='Plan',description='',labels=['D1'])]
    b.figures=[Figure(title='Combined views',description='',labels=['H±2','D1'])]
    assert not compare_source(a,b)
    b.figures[0].labels=['H±20','D1']
    assert compare_source(a,b)


def test_direct_grid_rejects_false_blank_observation(tmp_path):
    """Fast unit: reviewer cannot erase a directly confirmed merge. Expected <1 s, max 5 s."""
    t=CatalogTools(pdf(tmp_path),tmp_path/'out',role='reviewer');t.render(1)
    t.grids[1]=[{'row_count':1,'column_count':1,'expanded_rows':[['80']],'unknown_cells':[], 'merged_cells':[]}]
    wrong=content();wrong.tables[0].rows[0].cells[0].value=None;wrong.tables[0].rows[0].cells[0].state='blank'
    refused=t.save(1,wrong.model_dump_json(),'observation')
    assert not refused['saved'] and refused['direct_grid_differences']


def test_exhausted_resume_retains_review_and_content(tmp_path):
    """Focused integration: resume must not replace a rejected useful draft with pending. Expected <1 s, max 5 s."""
    from catalog_json.pipeline import convert
    class Workers:
        extractor=reviewer='test-model';contract_digest='test-contract';calls=[]
        async def run(self,role,source,out,page,attempt,feedback=False):
            self.calls.append(role)
            if role=='reviewer':
                r=review();r.verdict='revise';r.checks[0].passed=False;return r
            return content()
    w=Workers();source=pdf(tmp_path);out=tmp_path/'out'
    first=asyncio.run(convert(source,out,[1],w,max_repairs=0));assert first['pages'][0]['status']=='needs_review'
    final=asyncio.run(convert(source,out,[1],w,max_repairs=0,resume=True))
    assert final['pages'][0]['status']=='needs_review' and final['pages'][0]['content']==first['pages'][0]['content']
    assert final['pages'][0]['review']==first['pages'][0]['review'] and w.calls==['extractor','reviewer']


def test_direct_grid_keeps_merge_when_crop_slightly_cuts_border():
    """Focused integration: approximate agent crops must not lose spanning cells. Expected <1 s, max 5 s."""
    from catalog_json.grid import table_grids
    with pymupdf.open() as d:
        p=d.new_page(width=300,height=300)
        p.draw_line((30,60),(270,60));p.draw_line((30,120),(150,120))
        p.draw_line((30,180),(150,180));p.draw_line((150,180),(270,180))
        p.insert_text((50,90),'row1');p.insert_text((50,150),'row2');p.insert_text((200,120),'100')
        result=table_grids(p,[.09,.19,.91,.59]);table=result['tables'][0]
        assert table['expanded_rows']==[['row1','100'],['row2','100']]


def test_reviewer_can_revise_observations_more_than_twice(tmp_path):
    """Fast unit: source-grounded revisions continue until correct. Expected <1 s, max 5 s."""
    t=CatalogTools(pdf(tmp_path),tmp_path/'out');t.render(1);t.direct(1);t.save(1,content().model_dump_json(),'page')
    r=CatalogTools(t.source,t.output,role='reviewer');r.render(1);r.save(1,content().model_dump_json(),'observation');r.read(1,'page')
    for i in range(5):
        r.render(1,[0,0,1,1]);assert r.correct_observation(1,content().model_dump_json(),f'Fresh inspection {i}')['saved']
    assert len(list(r.work(1).glob('source-observation-0-revision-*.json')))==5


def test_default_repairs_continue_until_accepted(tmp_path):
    """Focused integration: quality loop has no arbitrary repair count; bounded by deadline. Expected <1 s, max 5 s."""
    from catalog_json.pipeline import convert
    class Workers:
        extractor=reviewer='test-model';contract_digest='test-contract'
        async def run(self,role,source,out,page,attempt,feedback=False):
            if role=='extractor':return content()
            r=review()
            if attempt<5:r.verdict='revise';r.checks[0].passed=False
            return r
    result=asyncio.run(convert(pdf(tmp_path),tmp_path/'out',[1],Workers()))
    assert result['complete_for_selection'] and len(result['pages'][0]['attempts'])==6


def test_repair_reads_latest_artifact_after_tenth_attempt(tmp_path):
    """Fast unit: unlimited repair counts must not select attempt 9 instead of 10. Expected <1 s, max 5 s."""
    from catalog_json.tools import write_json
    t=CatalogTools(pdf(tmp_path),tmp_path/'out',attempt=11)
    for number in (9,10):
        d=content();d.title=str(number);write_json(t.work(1)/f'candidate-{number}.json',d.model_dump())
    assert t.read(1,'page')['title']=='10'


def test_external_audit_is_source_bound_and_keeps_review_blind(tmp_path):
    """Fast unit: independent leads cannot cross source or bypass blind review; <1s/max5s."""
    import hashlib,pytest
    from catalog_json.tools import write_json
    source=pdf(tmp_path);out=tmp_path/'output'
    t=CatalogTools(source,out,page=1)
    write_json(t.work(1)/'external-audit.json',{'pdf_page':1,'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'findings':[{'path':'source_notes','observation':'Check tiny source text'}]})
    assert t.external_audit(1)['present']
    r=CatalogTools(source,out,role='reviewer',page=1)
    with pytest.raises(ValueError,match='independent'):r.external_audit(1)
    r.render(1);assert r.save(1,content().model_dump_json(),'observation')['saved']
    assert r.external_audit(1)['present']
    write_json(t.work(1)/'external-audit.json',{'pdf_page':2,'source_sha256':'wrong','findings':[]})
    with pytest.raises(ValueError,match='another source/page'):t.external_audit(1)


def test_installed_package_does_not_claim_ancestor_checkout_revision(tmp_path,monkeypatch):
    """Fast unit: installed/runtime metadata must not inherit another checkout; expected1s/max5s."""
    from catalog_json import provenance
    parent=tmp_path/'checkout';(parent/'.git').mkdir(parents=True)
    package=parent/'environment/lib/catalog_json';package.mkdir(parents=True)
    module=package/'provenance.py';module.write_text('# installed artifact')
    monkeypatch.setattr(provenance,'__file__',str(module))
    def forbidden(*args,**kwargs):raise AssertionError('Do not discover Git in package ancestors')
    monkeypatch.setattr(provenance.subprocess,'run',forbidden)
    result=provenance.program_version()
    assert result['git_revision'] is None and result['dirty'] is None


def test_cell_wrap_is_layout_without_relaxing_product_identifiers():
    from catalog_json.audit import same_value
    assert same_value('АС 70/72;\nАС 120/19; АС 120/27','АС 70/72; АС 120/19; АС 120/27',identity=True)
    assert not same_value('АС 70/72','АС 70/27',identity=True)
    assert not same_value('12 34','1234',identity=True)
    assert not same_value('СЗ','C3',identity=True)
    assert not same_value('1,0','1.0',identity=True)
    assert not same_value('М10-1','М10 1',identity=True)


def test_native_review_helpers_exist_in_initial_generation_snapshot(tmp_path):
    """Focused integration: prevent native enable-loop; expected2s/max20s; no LLM."""
    source=pdf(tmp_path)
    async def run():
        changed=asyncio.Event()
        def data(result):
            assert not result.is_error,result
            return result.structured_content if result.structured_content is not None else json.loads(next(x.text for x in result.content if getattr(x,'type',None)=='text'))
        async def message_handler(message):
            if getattr(getattr(message,'root',message),'method',None)=='notifications/tools/list_changed':changed.set()
        params=StdioServerParameters(command=sys.executable,args=['-m','catalog_json.tools','--source',str(source),'--output',str(tmp_path/'out'),'--role','reviewer','--page','1'])
        async with stdio_client(params) as (reader,writer):
            async with ClientSession(reader,writer,message_handler=message_handler) as client:
                await client.initialize();initial=await client.list_tools();names={t.name for t in initial.tools}
                assert {'view_page','view_region','save_review','enable_helpers','inspect_drawing_geometry',
                        'inspect_dimension_review_targets','record_dimension_reviews','read_source_table',
                        'read_source_layout','trace_dimension_candidates'}<=names
                assert 'open_catalog' not in names and 'read_review' not in names
                helper=next(t for t in initial.tools if t.name=='enable_helpers');assert 'ctx' not in helper.input_schema['properties']
                bad=await client.call_tool('enable_helpers',{'group':'invalid'});assert not data(bad)['enabled']
                result=await client.call_tool('enable_helpers',{'group':'labels'});assert data(result)['enabled']
                await asyncio.wait_for(changed.wait(),2)
                expanded={t.name for t in (await client.list_tools()).tools}
                assert expanded==names
                await client.call_tool('view_page',{'page':1})
                # Optional visibility never allows skipping blind observation and candidate inspection.
                refused=await client.call_tool('save_review',{'page':1,'json_text':review().model_dump_json()})
                assert not data(refused)['saved']
                await client.call_tool('enable_helpers',{'group':'drawings'})
                expanded={t.name for t in (await client.list_tools()).tools};assert expanded==names
    asyncio.run(asyncio.wait_for(run(),20))
