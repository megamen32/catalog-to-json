"""Focused integration: the delivered Excel reconstructs exact canonical JSON.

Purpose: catch type loss, Excel formula execution and 32K-cell truncation.
Expected 3s; emergency wall 300s. Fixtures only, no model/provider calls.
"""
import importlib.util
import json
from pathlib import Path
import zipfile

from openpyxl import load_workbook


MODULE = Path(__file__).resolve().parents[1] / "ops/catalogue_delivery/canonical_export.py"
spec = importlib.util.spec_from_file_location("canonical_export_test_subject", MODULE)
exporter = importlib.util.module_from_spec(spec)
spec.loader.exec_module(exporter)


def record(index, value, kind="model_parameter"):
    return {
        "id": f"char-{index}", "catalog": "test", "page": 1, "kind": kind,
        "model": "=MODEL", "parameter": "Диаметр", "part": None,
        "condition": "", "value": value, "unit": "мм", "state": "value",
        "source_fact_ids": [f"O-{index}"], "producer_ref": "original",
    }


def fixture(records):
    return {
        "schema": exporter.SCHEMA,
        "sources": [{"id": "test", "filename": "test.pdf", "page_count": 1}],
        "pages": [{"catalog": "test", "page": 1}],
        "producers": {"original": {"version": "fixture", "provider": "none"}},
        "metadata": {"review_scope": "fixture, not a catalogue review"},
        "records": records,
        "unresolved": [{"catalog": "test", "page": 1, "fact": {"value": None}}],
    }


def reconstruct(workbook_path):
    """Independent readback: no exporter serialization/decoding helpers used."""
    book = load_workbook(workbook_path, read_only=True, data_only=False)
    chunks = {}
    for sheet, row, record_id, field, index, count, value in book["Длинные значения"].iter_rows(min_row=2, values_only=True):
        chunks.setdefault((sheet, row, field), []).append((index, count, value))

    def complete(sheet, row_number, field, cell):
        parts = chunks.get((sheet, row_number, field))
        if parts is None:
            return cell
        parts.sort()
        assert [part[0] for part in parts] == list(range(len(parts)))
        assert {part[1] for part in parts} == {len(parts)}
        return "".join(part[2] for part in parts)

    records = []
    for name in ("Характеристики", "Чертежи", "Описание"):
        rows = book[name].iter_rows(values_only=False)
        headers = [cell.value for cell in next(rows)]
        for row_number, cells in enumerate(rows, 2):
            assert all(cell.data_type != "f" for cell in cells)
            values = dict(zip(headers, (cell.value for cell in cells)))
            for field in exporter.FIELDS:
                values[field] = complete(name, row_number, field, values[field])
            records.append({field: values[field] if field in ("id", "catalog", "page", "kind")
                            else json.loads(values[field]) for field in exporter.FIELDS})
    envelope = {}
    for row_number, (key, value) in enumerate(book["Метаданные"].iter_rows(min_row=2, values_only=True), 2):
        envelope[key] = json.loads(complete("Метаданные", row_number, "value", value))
    envelope["records"] = sorted(records, key=lambda item: item["id"])
    book.close()
    return envelope


def test_current_open_characteristics_are_visible_not_legacy_wrappers(tmp_path):
    records = [record(i, None) for i in range(25)]
    for item in records[:22]:
        item.update(catalog="small", state="source_conflict")
    for item in records[22:24]:
        item.update(catalog="small", page=218, model="SCT 6-14", state="unknown")
    data = fixture(records)
    data["unresolved"] = []  # All historical wrappers closed; engineering issues remain.
    source = tmp_path / "input.json"
    source.write_text(json.dumps(data), encoding="utf-8")
    result = exporter.export(source, tmp_path / "delivery")
    book = load_workbook(result["catalogue.xlsx"], read_only=False)
    issues = list(book["Неоднозначные"].values)
    assert issues[0] == ("Каталог", "Страница", "Модель", "Параметр", "state", "id")
    assert len(issues[1:]) == 24
    assert sum(row[4] == "source_conflict" for row in issues[1:]) == 22
    assert sum(row[4] == "unknown" for row in issues[1:]) == 2
    assert all(row[0] == "Малый каталог" for row in issues[1:])
    assert {row[5] for row in issues[1:]} == {r["id"] for r in records[:24]}
    assert book["Характеристики"].auto_filter.ref == "A1:N26"
    assert book["Неоднозначные"].auto_filter.ref == "A1:F25"
    book.close()
    manifest = result["manifest"]
    assert (manifest["source_conflict_count"], manifest["unknown_count"],
            manifest["unresolved_count"], manifest["legacy_unresolved_count"]) == (22, 2, 24, 0)
    readme = Path(result["README.txt"]).read_text()
    assert "Открытых противоречий (source_conflict): 22." in readme
    assert "Неизвестных значений (unknown): 2." in readme
    assert "Всего нерешённых характеристик на листе «Неоднозначные»: 24." in readme
    assert reconstruct(result["catalogue.xlsx"])["records"] == sorted(records, key=lambda r: r["id"])
    assert Path(result["catalogue.json"]).read_bytes() == source.read_bytes()


def test_typed_values_source_provenance_and_formula_are_preserved(tmp_path):
    values = [None, "", "—", 0, 0.0, False, "=1+1", "+SUM(A1)", "01",
              {"nominal": 25, "tolerance": [0, -1.5]}, ["А", None, 3],
              "первая\r\nвторая", "control\x00\ufffe", 10 ** 40]
    source_data = fixture([record(index, value, ("model_parameter", "figure", "text")[index % 3])
                           for index, value in enumerate(values)])
    source = tmp_path / "input.json"
    source.write_text(json.dumps(source_data, ensure_ascii=False), encoding="utf-8")
    result = exporter.export(source, tmp_path / "delivery")
    expected = {**source_data, "records": sorted(source_data["records"], key=lambda item: item["id"])}
    readback = reconstruct(result["catalogue.xlsx"])
    assert readback == expected
    # Python considers 0 == 0.0 == False; explicitly prove their types too.
    assert {item["id"]: type(item["value"]) for item in readback["records"]} == {
        item["id"]: type(item["value"]) for item in source_data["records"]}
    assert Path(result["catalogue.json"]).read_bytes() == source.read_bytes()
    manifest = json.loads(Path(result["manifest.json"]).read_text())
    for name, entry in manifest["files"].items():
        path = tmp_path / "delivery" / name
        assert entry["sha256"] == exporter.sha256(path)
        assert entry["bytes"] == path.stat().st_size
    with zipfile.ZipFile(result["catalogues-unified.zip"]) as archive:
        assert set(archive.namelist()) == {"catalogue.json", "catalogue.xlsx", "README.txt", "manifest.json"}
        assert archive.testzip() is None
        for name in archive.namelist():
            assert archive.read(name) == (tmp_path / "delivery" / name).read_bytes()
    assert not list((tmp_path / "delivery").glob(".xlsx-*"))


def test_long_unicode_text_and_envelope_reconstruct_without_truncation(tmp_path):
    text = "🛠АБ\n" * 18_000
    item = record(0, text, "figure")
    item["condition"] = text
    data = fixture([item])
    data["metadata"]["detail"] = text
    data["unresolved"][0]["fact"]["value"] = text
    source = tmp_path / "input.json"
    source.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    result = exporter.export(source, tmp_path / "delivery")
    assert reconstruct(result["catalogue.xlsx"]) == data
    book = load_workbook(result["catalogue.xlsx"], read_only=True)
    chunks = list(book["Длинные значения"].iter_rows(min_row=2, values_only=True))
    assert len(chunks) >= 12
    assert all(len(row[-1].encode("utf-16-le")) // 2 <= 30_000 for row in chunks)
    book.close()


def test_existing_deliverable_is_not_overwritten(tmp_path):
    source = tmp_path / "input.json"
    source.write_text(json.dumps(fixture([record(0, 1)])))
    output = tmp_path / "delivery"
    output.mkdir()
    existing = output / "catalogue.json"
    existing.write_text("previous accepted file")
    import pytest
    with pytest.raises(FileExistsError):
        exporter.export(source, output)
    assert existing.read_text() == "previous accepted file"


def test_actual_merge_builder_exports_its_output_without_schema_translation(tmp_path, monkeypatch):
    """Focused integration: exercise the real build→export contract boundary.

    Synthetic source receipts cover the production 980-page shape. Only archive
    identity/provenance lookup and builder Git metadata are fixtures; build,
    output serialization, export and independent workbook readback are real.
    """
    merge_spec = importlib.util.spec_from_file_location("canonical_merge_integration", MODULE.with_name("canonical_merge.py"))
    merge = importlib.util.module_from_spec(merge_spec)
    merge_spec.loader.exec_module(merge)
    pages, packets = tmp_path / "pages", tmp_path / "packets"
    pages.mkdir()
    packets.mkdir()
    archive = tmp_path / "immutable-input.zip"
    archive.write_bytes(b"synthetic identity, page producer lookup is injected")
    monkeypatch.setattr(merge, "INPUT_SHA", exporter.sha256(archive))
    monkeypatch.setattr(merge, "program_identity", lambda: {"name": "fixture-builder", "version": "test"})
    producers = {}
    for catalog_number, (catalog, page_count) in enumerate(merge.COUNTS.items(), 1):
        source_sha = f"{catalog_number:064x}"
        for page_number in range(1, page_count + 1):
            facts, groups = [], []
            if catalog_number == 1 and page_number == 1:
                facts = [{
                    "fact_id": "O-real-seam-fixture", "side": "ours", "catalog": catalog,
                    "physical_pages": [page_number], "kind": "model_parameter",
                    "owner": "FIXTURE-1", "name": "Диаметр", "value": "0",
                    "unit": "мм", "condition": None, "pointer": "/saved/source/cell",
                    "metadata": {},
                }]
                groups = [{"keep": "ours", "relation": "ours_only", "fact_indices": [0],
                           "reason": "Synthetic retained-source fixture"}]
            receipt = {"source_sha256": source_sha, "work_sha256": "b" * 64,
                       "identity": {"runtime": "fixture", "program_version": "test"}, "groups": groups}
            filename = f"{catalog}-{page_number:04}.json"
            (pages / filename).write_text(json.dumps({"catalog": catalog, "page": page_number,
                                                     "facts": facts, "receipt": receipt}))
            (packets / filename).write_text('{"contexts":{}}')
            producers[(source_sha, page_number)] = {"source_filename": catalog + ".pdf",
                                                   "candidate_producer": {"version": "fixture"},
                                                   "extractor": "no provider"}
    monkeypatch.setattr(merge, "load_page_producers", lambda unused_archive: producers)
    amendments = tmp_path / "amendments.json"
    amendments.write_text('{"actions":[]}')
    unresolved = tmp_path / "unresolved.json"
    unresolved.write_text(json.dumps({"facts": [
        {"catalog": "opn", "physical_page": 1, "fact": {"fact_id": f"unresolved-{index}"}}
        for index in range(42)
    ]}))
    built = tmp_path / "built"
    report = merge.build(pages, packets, archive, amendments, unresolved, built)
    assert report["records"] == 1 and report["pages"] == 980 and report["conflicts"] == 0
    # No pre-export schema repair: this failed when build wrote schema_version
    # while the exporter correctly required the agreed canonical schema key.
    paths = exporter.export(built / "catalogue.json", tmp_path / "delivery")
    assert reconstruct(paths["catalogue.xlsx"]) == json.loads((built / "catalogue.json").read_text())
