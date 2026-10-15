"""Fast unit tests for independent readback, expected <5s, emergency 50s.

Detect type collapse, chunk loss, hidden physical conflicts, missing pages and
changed retained unresolved facts. Fixtures deliberately do not use exporter.
"""
import importlib.util
import json
from pathlib import Path
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

SPEC = importlib.util.spec_from_file_location(
    "verify_canonical_under_test",
    Path(__file__).parents[1] / "ops/catalogue_delivery/verify_canonical.py")
verify_module = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(verify_module)


def j(value):
    return json.dumps(value, ensure_ascii=False, allow_nan=False, separators=(",", ":"))


def make_files(tmp_path, values=None):
    if values is None:
        values = [None, "", "—", 0, False, "0", {"a": [0, None, ""]}, "=SUM(1,2)"]
    records = []
    for n, value in enumerate(values):
        records.append(dict(zip(verify_module.CORE, [
            f"r{n}", "test", 1, "model_parameter", f"model{n}", "parameter",
            None, "", value, "mm", "source_value", [f"O-{n}"], "p1"])))
    unresolved = [{"catalog": "test", "physical_page": 1, "fact": {"id": str(n)}} for n in range(42)]
    doc = {"schema": "canonical-catalogue/1", "sources": [{"id": "test", "filename": "source.pdf", "sha256": "a"*64, "page_count": 2}],
           "pages": [{"catalog": "test", "page": 1}, {"catalog": "test", "page": 2}],
           "records": records, "producers": {"p1": {"actual_version": "1.2.3"}},
           "unresolved": unresolved, "metadata": {"input_archive_sha256": "b"*64, "selected_table_source_fact_ids": ["O-0"]}}
    json_path = tmp_path / "catalogue.json"
    json_path.write_text(j(doc), encoding="utf-8")
    reference = tmp_path / "unresolved.json"
    reference.write_text(j({"facts": unresolved, "source_archive_sha256": "b"*64}), encoding="utf-8")
    book = Workbook()
    book.remove(book.active)
    for name in verify_module.DATA_SHEETS:
        sheet = book.create_sheet(name)
        sheet.append(list(verify_module.CORE) + ["value_display"])
    for record in records:
        sheet = book[verify_module._sheet_for(record)]
        sheet.append([record[field] if field in verify_module.LITERAL else j(record[field]) for field in verify_module.CORE] + ["display"])
    sheet = book.create_sheet("Неоднозначные")
    sheet.append(["index", "value"])
    for n, value in enumerate(unresolved):
        sheet.append([n, j(value)])
    sheet = book.create_sheet("Метаданные")
    sheet.append(["key", "value"])
    for key, value in doc.items():
        if key not in {"records", "unresolved"}:
            sheet.append([key, j(value)])
    xlsx = tmp_path / "catalogue.xlsx"
    book.save(xlsx)
    return doc, json_path, xlsx, reference


def check(files):
    _, json_path, xlsx, reference = files
    return verify_module.verify(json_path, xlsx, unresolved_reference=reference,
                                expected_table_ids=["O-0"], expected_page_counts={"test": 2})


def codes(result):
    return {failure["code"] for failure in result["failures"]}


def test_readback_preserves_types_and_empty_page(tmp_path):
    result = check(make_files(tmp_path))
    assert result["ok"], result
    assert result["coverage"]["pages"] == 2
    assert result["coverage"]["pages_with_records"] == 1


@pytest.mark.parametrize("replacement", ['""', '"0"', "false", '"—"'])
def test_null_does_not_collapse_into_other_values(tmp_path, replacement):
    files = make_files(tmp_path)
    book = load_workbook(files[2])
    book["Характеристики"].cell(2, 9, replacement)
    book.save(files[2])
    result = check(files)
    assert "xlsx_field_mismatch" in codes(result)


def test_formula_not_accepted_as_value(tmp_path):
    files = make_files(tmp_path)
    book = load_workbook(files[2])
    book["Характеристики"].cell(2, 9, "=1+2")
    book.save(files[2])
    assert "xlsx_open_or_decode" in codes(check(files))


def test_same_physical_characteristic_across_pages_is_conflict(tmp_path):
    files = make_files(tmp_path, values=[1, 2])
    doc = files[0]
    doc["records"][1]["model"] = doc["records"][0]["model"]
    doc["records"][1]["page"] = 2
    files[1].write_text(j(doc), encoding="utf-8")
    assert "physical_value_conflict" in codes(check(files))


def test_missing_explicit_page_fails_even_when_records_open(tmp_path):
    files = make_files(tmp_path)
    files[0]["pages"].pop()
    files[1].write_text(j(files[0]), encoding="utf-8")
    assert "page_manifest_coverage" in codes(check(files))


def test_unresolved_content_cannot_be_changed_at_same_count(tmp_path):
    files = make_files(tmp_path)
    files[0]["unresolved"][0]["fact"]["id"] = "replacement"
    files[1].write_text(j(files[0]), encoding="utf-8")
    assert "retained_unresolved_mismatch" in codes(check(files))


def test_chunk_reconstruction_and_missing_chunk(tmp_path):
    files = make_files(tmp_path, values=["long " * 10])
    book = load_workbook(files[2])
    book["Характеристики"].cell(2, 9, verify_module.LONG_MARKER)
    sheet = book.create_sheet("Длинные значения")
    sheet.append(["sheet", "row", "record_id", "field", "chunk_index", "chunk_count", "chunk"])
    value = j(files[0]["records"][0]["value"])
    sheet.append(["Характеристики", 2, "r0", "value", 0, 2, value[:20]])
    sheet.append(["Характеристики", 2, "r0", "value", 1, 2, value[20:]])
    book.save(files[2])
    assert check(files)["ok"]
    sheet.delete_rows(3)
    book.save(files[2])
    assert "xlsx_open_or_decode" in codes(check(files))


def test_whole_table_missing_fact_fails(tmp_path):
    files = make_files(tmp_path)
    result = verify_module.verify(files[1], files[2], unresolved_reference=files[3],
                                  expected_table_ids=["missing"], expected_page_counts={"test": 2})
    assert "selected_table_facts_missing" in codes(result)


def test_readback_does_not_trust_display_column(tmp_path):
    files = make_files(tmp_path)
    book = load_workbook(files[2])
    book["Характеристики"].cell(2, 14, "arbitrary display text")
    book.save(files[2])
    assert check(files)["ok"]


def test_manifest_archive_actual_bytes_are_verified(tmp_path):
    files = make_files(tmp_path)
    readme = tmp_path / "README.txt"
    readme.write_text("Catalogue", encoding="utf-8")
    manifest = tmp_path / "manifest.json"
    selected_paths = [files[1], files[2], readme]
    manifest.write_text(j({"files": {path.name: {"sha256": verify_module.digest(path),
                                              "bytes": path.stat().st_size}
                                     for path in selected_paths}}), encoding="utf-8")
    archive = tmp_path / "delivery.zip"
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as zipped:
        for path in selected_paths + [manifest]:
            zipped.write(path, path.name)
    kwargs = {"unresolved_reference": files[3], "expected_table_ids": ["O-0"],
              "expected_page_counts": {"test": 2}, "manifest": manifest, "archive": archive}
    assert verify_module.verify(files[1], files[2], **kwargs)["ok"]
    readme.write_text("Different bytes", encoding="utf-8")
    result = verify_module.verify(files[1], files[2], **kwargs)
    assert "manifest_hash_or_size" in codes(result)
    assert "archive_file_bytes" in codes(result)


def test_absent_xlsx_returns_failure_report_not_exception(tmp_path):
    files = make_files(tmp_path)
    files[2].unlink()
    result = check(files)
    assert not result["ok"]
    assert "xlsx_open_or_decode" in codes(result)


def test_extra_canonical_fields_cannot_silently_disappear(tmp_path):
    files = make_files(tmp_path)
    files[0]["records"][0]["additional_evidence"] = "must not be dropped"
    files[1].write_text(j(files[0]), encoding="utf-8")
    assert "unrepresented_record_fields" in codes(check(files))
