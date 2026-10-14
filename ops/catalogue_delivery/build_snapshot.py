"""User delivery: bounded snapshot, expected90s/max900s, RAM3GiB, one CPU worker, temp512MiB; no model calls."""
from pathlib import Path
import json,hashlib,datetime,collections,resource,time,shutil,zipfile,os,gc
resource.setrlimit(resource.RLIMIT_AS,(3*2**30,3*2**30));resource.setrlimit(resource.RLIMIT_CPU,(600,600));resource.setrlimit(resource.RLIMIT_FSIZE,(256*2**20,256*2**20))
from catalog_json.schema import PageContent
from catalog_json.export import to_xlsx
from openpyxl import load_workbook
r=Path.cwd();started=time.time();base=r/'.tmp/full-digitization/staged-best-v1.4.8';refs=r/'examples/manual-review/2026-10-14';stamp=datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%dT%H%M%SZ');out=r/'examples/catalogs'/('all-pages-'+stamp);out.mkdir();pack=r/'.tmp/root-catalogue-pages-20261014'/('package-'+stamp);pack.mkdir();manual={('armature',27),('isolators',47),('innovations',168)};catalogs=[];snapshots=[];total_cells=0;accepted=0
from completed_pages import completed,sha
from catalog_json.manual_completion import independent_completion_update
# Reconcile latest immutable authored revisions, not lexicographic index order.
tail={(x['catalog'],x['physical_page']):(Path(x['content']),x['proof']) for x in completed(r)}
independent_snapshot=json.loads((r/'.tmp/root-tail-review-20261014/independent-status.json').read_text())
passes={(x['catalog'],x['physical_page']):x for x in independent_snapshot['pages'] if x['verdict']=='pass'}
all_reviews={(x['catalog'],x['physical_page']):x for x in independent_snapshot['pages']}
manual_verified=0;native_queue_accepted=0
def public_metadata(v):
 if isinstance(v,dict):return {k:public_metadata(x) for k,x in v.items()}
 if isinstance(v,list):return [public_metadata(x) for x in v]
 if isinstance(v,str):
  import re
  return re.sub(r'/(?:home|Users)/[^\s,;]+',lambda m:Path(m.group()).name,v)
 return v
manual_count=0
for cid in ['opn','isolators','small','innovations','armature']:
 f=base/cid/'catalog.json';raw=f.read_bytes();old=json.loads(raw);native_queue_accepted+=sum(x['status']=='accepted' for x in old['pages']);compact={'id':cid,'source':old['source'],'models':old['models'],'complete_for_selection':False,'catalog_complete':False,'pages':[]};snapshots.append({'catalog':cid,'checkpoint_sha256':hashlib.sha256(raw).hexdigest(),'checkpoint_mtime_unix':f.stat().st_mtime,'snapshot_read_unix':time.time()});del raw
 for p in old['pages']:
  pg=p['pdf_page'];w=base/cid/f'page-{pg:04d}';content=PageContent.model_validate(p['content']);producer=p['candidate_producer'];origin=w/f"candidate-{producer['attempt']}.json";origin_raw=origin.read_bytes();assert hashlib.sha256(origin_raw).hexdigest()==producer['artifact_sha256'],(cid,pg,'candidate bytes changed');assert PageContent.model_validate_json(origin_raw).model_dump()==content.model_dump(),(cid,pg,'content does not match actual producer');assert producer['status']=='receipt_verified' and producer['source_sha256']==old['source']['sha256'] and producer['physical_page']==pg
  rec={k:p[k] for k in ['pdf_page','status','source_sha256','candidate_sha256','candidate_producer','review_producer','review_sha256','review','role_contracts','extractor','reviewer','elapsed_seconds'] if k in p};rec['content']=content.model_dump();rec['source_filename']=old['source']['filename'];rec['source_checked_by']=p.get('review_producer',{}).get('session_id') if isinstance(p.get('review_producer'),dict) and p['review_producer'].get('status')=='manual_source_review_verified' else None;rec['reviewed_current_content']=p['status']=='accepted'
  if rec.get('review_producer'):
   rv=rec['review_producer'];rf=w/'root-source-review.json' if rv['status']=='manual_source_review_verified' else w/f"review-{rv['attempt']}.json";assert hashlib.sha256(rf.read_bytes()).hexdigest()==rv['artifact_sha256'];assert rv['source_sha256']==old['source']['sha256'] and rv['physical_page']==pg
  if (cid,pg) in tail and (tail[(cid,pg)][1].get('corrections') or (cid,pg) in passes):
   mf,proof=tail[(cid,pg)];mr=mf.read_bytes();author=proof['reviewer'];rec={'pdf_page':pg,'source_filename':old['source']['filename'],'source_sha256':old['source']['sha256'],'status':'root_source_checked_awaiting_independent_review','candidate_sha256':proof['output']['sha256'],'candidate_producer':{'status':'manual_source_checked','harness':'codex','session_id':author['session_id'],'model':author.get('model_id') or author.get('model') or 'Native Codex session; exact provider model identifier not recorded','program':{'package_version':author['version'],'source_code_sha256':None},'method':'Full original-page visual check by independently assigned Codex author','artifact_sha256':proof['output']['sha256'],'source_sha256':old['source']['sha256'],'physical_page':pg,'authored_at':proof['completed_at'],'native_basis_sha256':proof['basis_candidate']['sha256']},'source_checked_by':author['session_id'],'reviewed_current_content':False,'review_producer':None,'content':PageContent.model_validate_json(mr).model_dump()};manual_count+=1
  elif (cid,pg) in manual and not rec['reviewed_current_content']:
   mf=refs/f'{cid}-page-{pg:04d}.json';prov=json.loads(mf.with_suffix('.provenance.json').read_text());mr=mf.read_bytes();assert hashlib.sha256(mr).hexdigest()==prov['output']['sha256'];assert prov['source']['sha256']==old['source']['sha256'] and prov['source']['physical_page']==pg
   manual_count+=1
   rec={'pdf_page':pg,'source_filename':old['source']['filename'],'source_sha256':old['source']['sha256'],'status':'root_source_checked_awaiting_independent_review','candidate_sha256':prov['output']['sha256'],'candidate_producer':{'status':'manual_source_checked','harness':'codex','session_id':prov['producer']['session_id'],'model':prov['producer']['model_id'],'program':{'package_version':prov['producer']['version'],'source_code_sha256':None},'method':prov['producer']['method'],'artifact_sha256':prov['output']['sha256'],'source_sha256':old['source']['sha256'],'physical_page':pg,'authored_at':prov['created_at']},'source_checked_by':'Codex, личная сверка оригинала','reviewed_current_content':False,'review_producer':None,'content':PageContent.model_validate_json(mr).model_dump()}
  elif rec['status']!='accepted':
   rec['reviewed_current_content']=False
   if not rec.get('review_producer'):rec['review_producer']=None
  # A manual correction receives a separate, real second-author source review.
  # Native provider receipts and immutable basis files are never rewritten.
  ip=passes.get((cid,pg))
  if ip and rec['candidate_producer']['status']=='manual_source_checked' and tail.get((cid,pg),(None,{}))[1].get('corrections'):
   completion=independent_completion_update(w,p)
   marker=(completion or {}).get('independent_revision_completion',{})
   if marker.get('reference_sha256')!=rec['candidate_sha256'] or marker.get('review_sha256')!=ip['review_sha256']:ip=None
  if rec['candidate_producer']['status']=='manual_source_checked' and ip and ip['content_sha256']==rec['candidate_sha256']:
   report_path=Path(ip['review']);assert sha(report_path)==ip['review_sha256'];report=json.loads(report_path.read_text())
   assert ip['source_sha256']==old['source']['sha256'] and ip['reviewer']!=rec['candidate_producer']['session_id']
   assert ip['content_sha256']==sha(Path(ip['content'])) and ip['author']==rec['candidate_producer']['session_id']
   assert report['verdict']=='pass' and report['full_original_image_review'] is True
   assert not report.get('issues',report.get('errors',[])) and not rec['content']['uncertainties']
   ver=report.get('version',report['schema_version']);rec['status']='manual_independently_source_verified';rec['reviewed_current_content']=True
   rec['review_producer']={'status':'manual_independent_source_review_verified','harness':'codex','session_id':ip['reviewer'],'model':'Native Codex agent; exact base model identity not exposed','program':{'package_version':ver,'source_code_sha256':None},'artifact_sha256':ip['review_sha256'],'reviewed_candidate_sha256':ip['content_sha256'],'source_sha256':ip['source_sha256'],'physical_page':pg,'provider_receipt':False,'method':'Different assigned author reviewed the full original page and exact frozen corrected JSON'}
   rec['review']=public_metadata(report);rec['review_sha256']=ip['review_sha256'];rec['source_checked_by']=[ip['author'],ip['reviewer']];manual_verified+=1
  # Retain an actual independent REVISE too: lack of acceptance is not lack of review.
  ir=all_reviews.get((cid,pg))
  if rec['candidate_producer']['status']=='manual_source_checked' and ir and ir['content_sha256']==rec['candidate_sha256']:
   report_path=Path(ir['review']);assert sha(report_path)==ir['review_sha256'];report=json.loads(report_path.read_text())
   assert ir['reviewer']!=rec['candidate_producer']['session_id'] and report['full_original_image_review'] is True
   rec['independent_source_review']={'verdict':ir['verdict'],'reviewer':ir['reviewer'],'review_sha256':ir['review_sha256'],'reviewed_candidate_sha256':ir['content_sha256'],'full_original_image_review':True,'provider_model_call':False,'reviewed_at':report.get('time',report.get('reviewed_at')),'issues':public_metadata(report.get('issues',report.get('errors',[]))),'unresolved_issues':public_metadata(report.get('unresolved_issues',report.get('uncertainties',[])))}
   if not rec['reviewed_current_content'] and ir['verdict']=='revise':
    rec['status']='source_quality_limited' if rec['content']['uncertainties'] and not rec['independent_source_review']['issues'] else 'independently_reviewed_needs_correction'
  # Preserve actual native timing as timing of the native basis, not manual elapsed time.
  if rec['candidate_producer']['status']=='manual_source_checked':rec['native_basis_elapsed_seconds']=p.get('elapsed_seconds')
  for metadata_key in [k for k in rec if k!='content']:rec[metadata_key]=public_metadata(rec[metadata_key])
  rec['content_canonical_sha256']=hashlib.sha256(json.dumps(rec['content'],ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest();compact['pages'].append(rec)
 compact['catalog_complete']=all(p['reviewed_current_content'] for p in compact['pages']);compact['complete_for_selection']=compact['catalog_complete'];compact['counts']={'pages':len(compact['pages']),'independently_accepted':sum(p['reviewed_current_content'] for p in compact['pages']),'root_source_checked_pending_review':sum(not p['reviewed_current_content'] and p['candidate_producer']['status']=='manual_source_checked' for p in compact['pages'])};accepted+=compact['counts']['independently_accepted'];total_cells+=sum(len(row['cells']) for p in compact['pages'] for t in p['content']['tables'] for row in t['rows']);catalogs.append(compact);del old
 x=pack/(cid+'.xlsx');to_xlsx({**compact,'pages':[{**v,'review_producer':v.get('review_producer') or {}} for v in compact['pages']]},x,allow_unverified=True);book=load_workbook(x);prov_sheet=book['Происхождение'];prov_sheet.cell(1,14,'Лично сверил исходник');prov_sheet.cell(1,15,'Независимо принят именно этот разбор')
 for row,p in enumerate(compact['pages'],2):
  prov_sheet.cell(row,14,', '.join(p['source_checked_by']) if isinstance(p['source_checked_by'],list) else p['source_checked_by']);prov_sheet.cell(row,15,'да' if p['reviewed_current_content'] else 'нет')
  if p['candidate_producer']['status']=='manual_source_checked':
   prov_sheet.cell(row,6,'Codex: личный ручной разбор');prov_sheet.cell(row,7,p['candidate_producer']['program']['package_version']);prov_sheet.cell(row,9,'Codex: '+p['review_producer']['session_id'] if p['reviewed_current_content'] else 'Ожидается независимая проверка');prov_sheet.cell(row,10,p['review_producer']['program']['package_version'] if p['reviewed_current_content'] else 'не выполнена')
  elif not p.get('review_producer'):
   prov_sheet.cell(row,9,'Проверка ещё не завершена');prov_sheet.cell(row,10,'не выполнена')
  if p.get('independent_source_review'):
   ir=p['independent_source_review'];prov_sheet.cell(1,16,'Выполненная независимая сверка');prov_sheet.cell(1,17,'Что мешает приёмке')
   prov_sheet.cell(row,16,ir['reviewer']+': '+ir['verdict'])
   if not p['reviewed_current_content']:
    prov_sheet.cell(row,9,'Codex: '+ir['reviewer']);prov_sheet.cell(row,10,'Независимая сверка выполнена, приёмка открыта')
    prov_sheet.cell(row,17,'; '.join(p['content']['uncertainties']) or json.dumps(ir['issues'],ensure_ascii=False))
 book.save(x);del book;gc.collect();print({'catalog':cid,**compact['counts'],'xlsx_bytes':x.stat().st_size},flush=True)
assert sum(len(c['pages']) for c in catalogs)==980
bundle={'schema_version':'catalogue-snapshot/3','created_at':datetime.datetime.now().astimezone().isoformat(),'snapshot_stage':'partial_all_pages','complete_for_selection':all(c['complete_for_selection'] for c in catalogs),'page_count':980,'independently_accepted_pages':accepted,'root_source_checked_pages_awaiting_independent_review':manual_count-manual_verified,'manual_pages_independently_source_verified':manual_verified,'native_queue_accepted_pages':native_queue_accepted,'coordinator_program_version':'1.5.24','exporter_program':{'version':'unified-catalogue-snapshot/5','source_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest()},'manual_producer_version':'Per-page actual native Codex author and source-review version','notes':['Включены все готовые разборы всех пяти каталогов, независимо от завершения проверки.','Ручные разборы root и трёх сабагентов Codex интегрированы в свои страницы, отдельных конкурирующих версий внутри архива нет.','accepted — приёмка штатной очередью; manual_independently_source_verified — отдельная проверка другим Codex-исполнителем по оригиналу и точному SHA. Остальные статусы не означают PASS.','Автор/версия/исходник/страница сохранены индивидуально. Старые модельные проверки не присвоены новым ручным разборам.','Координатор1.5.24 не является автором всех страниц: фактическая версия каждого извлекателя/проверяющего указана в producer.'],'catalogs':catalogs}
bundle['pages_requiring_source_clarification']=[{'catalog':c['id'],'physical_page':p['pdf_page'],'source_filename':p['source_filename'],'source_sha256':p['source_sha256'],'status':p['status'],'uncertainties':p['content']['uncertainties'],'independent_source_review':p.get('independent_source_review')} for c in catalogs for p in c['pages'] if not p['reviewed_current_content']]
(pack/'all-catalogues-980-pages.json').write_text(json.dumps(bundle,ensure_ascii=False,indent=2)+'\n');(pack/'README.txt').write_text('Снимок всех пяти каталогов — '+bundle['created_at']+'\n\nВсе 980 физических страниц включены. Независимо приняты: '+str(accepted)+'. Это промежуточная выдача, не весь каталог проверен.\n\nall-catalogues-980-pages.json — единый JSON с пятью каталогами.\nopn.xlsx, isolators.xlsx, small.xlsx, innovations.xlsx, armature.xlsx — Excel по тем же данным.\nВ Excel лист «Происхождение» показывает физическую/печатную страницу, исходный PDF и SHA256, автора/версию и проверяющего/версию, статус и принятие именно текущего разбора.\n\nРучные разборы root и трёх сабагентов по mod4,а также прежние27/47/168 уже встроены в соответствующие каталоги. Их автор и отдельный проверяющий указаны по каждой странице. manual_independently_source_verified означает проверку другим исполнителем по полному оригиналу; это не GLM receipt. Для страниц с замечаниями сохранены отдельный independent_source_review, имя проверяющего и причина открытой приёмки; выполненная сверка с REVISE не выдаётся за PASS.\n\nВремя и токены: это снимок существующих выходов, новых запросов моделям при сборке не было. Исторические транскрипты и секреты в архив не включены.\n',encoding='utf-8')
with (pack/'README.txt').open('a') as f:
 f.write('\nСтраницы с открытыми замечаниями (физические номера PDF):\n')
 for item in bundle['pages_requiring_source_clarification']:f.write(item['catalog']+' '+str(item['physical_page'])+': '+'; '.join(item['uncertainties'])+'\n')
files={p.name:{'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'bytes':p.stat().st_size} for p in pack.iterdir()};manifest={'created_at':bundle['created_at'],'page_count':980,'independently_accepted_pages':accepted,'manual_pages_integrated':manual_count,'manual_independently_source_verified':manual_verified,'native_queue_accepted_pages':native_queue_accepted,'counts':{c['id']:c['counts'] for c in catalogs},'files':files,'snapshot_sources':snapshots,'budget':{'expected_seconds':90,'maximum_seconds':900,'ram_hard_bytes':3*2**30,'cpu_workers':1,'temporary_disk_max_bytes':512*2**20,'model_calls':0}};(pack/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n');assert sum(p.stat().st_size for p in pack.iterdir())<512*2**20
zip_path=out/'all-catalogues-980-pages.zip'
with zipfile.ZipFile(zip_path,'x',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as z:
 for p in sorted(pack.iterdir()):z.write(p,p.name)
verify=r/'.tmp/root-catalogue-pages-20261014'/('unpacked-'+stamp);verify.mkdir()
with zipfile.ZipFile(zip_path) as z:
 assert z.testzip() is None
 assert all(Path(n).name==n for n in z.namelist())
 z.extractall(verify)
loaded=json.loads((verify/'all-catalogues-980-pages.json').read_text());assert len(loaded['catalogs'])==5 and sum(len(c['pages']) for c in loaded['catalogs'])==980;checked_cells=0;checked_pages=0;checked_figures=0
for name,entry in files.items():assert hashlib.sha256((verify/name).read_bytes()).hexdigest()==entry['sha256']
for c in loaded['catalogs']:
 book=load_workbook(verify/(c['id']+'.xlsx'),read_only=True,data_only=False);expected_cells=[];text_rows=1
 for p in c['pages']:
  content=PageContent.model_validate(p['content']);assert hashlib.sha256(json.dumps(p['content'],ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()==p['content_canonical_sha256'];checked_pages+=1;checked_figures+=len(content.figures);text_rows+=1+len(content.text_blocks)+len(content.figures)+len(content.source_notes)+len(content.uncertainties)
  for i,t in enumerate(content.tables,1):
   sheet=book[f'Стр {p["pdf_page"]} таб {i}'];actual=list(sheet.iter_rows(values_only=True));assert len(actual)==len(t.rows)+1 and len(actual[0])==len(t.columns),(c['id'],p['pdf_page'],'shape')
   cols={col.id:col for col in t.columns}
   for actual_row,source in zip(actual[1:],t.rows):
    cells={cell.column_id:cell for cell in source.cells};want=tuple(None if cells[col.id].value=='' else cells[col.id].value for col in t.columns);assert tuple(actual_row)==want,(c['id'],p['pdf_page'],'cell roundtrip');checked_cells+=len(want)
    for cell in source.cells:col=cols[cell.column_id];expected_cells.append((p['pdf_page'],t.title,len(expected_cells),col.label,cell.value,col.unit,cell.state,col.group,col.qualifier))
  if p.get('independent_source_review'):
   ir=p['independent_source_review'];assert book['Происхождение'].cell(p['pdf_page']+1,16).value==ir['reviewer']+': '+ir['verdict']
   if not p['reviewed_current_content']:assert book['Происхождение'].cell(p['pdf_page']+1,17).value==('; '.join(p['content']['uncertainties']) or json.dumps(ir['issues'],ensure_ascii=False))
  if p['candidate_producer']['status']=='manual_source_checked':
   if p['reviewed_current_content']:assert p['review_producer']['status']=='manual_independent_source_review_verified' and p['review_producer']['reviewed_candidate_sha256']==p['candidate_sha256'] and p['review_producer']['session_id']!=p['candidate_producer']['session_id']
   else:assert p['review_producer'] is None
 assert book['Данные'].max_row==len(expected_cells)+1;assert book['Текст и чертежи'].max_row==text_rows;assert book['Происхождение'].max_row==len(c['pages'])+1;book.close();gc.collect()
assert checked_cells==total_cells
receipt={'zip':str(zip_path.relative_to(r)),'zip_sha256':hashlib.sha256(zip_path.read_bytes()).hexdigest(),'zip_bytes':zip_path.stat().st_size,'archive_CRC':'passed','extracted_files':len(list(verify.iterdir())),'json_pages_validated':checked_pages,'excel_workbooks_reopened':5,'independent_review_attribution_roundtrip':True,'excel_data_cells_exact_roundtrip':checked_cells,'figure_descriptions_in_json':checked_figures,'manual_pages_integrated':manual_count,'manual_independently_source_verified':manual_verified,'native_queue_accepted_pages':native_queue_accepted,'independent_accepted':accepted,'all_pages_independently_reviewed':False,'model_calls':0,'elapsed_seconds':round(time.time()-started,3),'peak_rss_bytes':resource.getrusage(resource.RUSAGE_SELF).ru_maxrss*1024};(out/'verification.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n');shutil.copy2(pack/'manifest.json',out/'manifest.json');(r/'.tmp/root-catalogue-pages-20261014/latest-delivery.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n');print(receipt,flush=True)
