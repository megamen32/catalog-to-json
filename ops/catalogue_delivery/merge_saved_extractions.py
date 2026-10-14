"""One portable catalogue snapshot from retained inputs and actual three-party votes.

No provider calls. Source inputs/receipts are never edited. A partial snapshot is
explicitly opt-in; full delivery additionally requires full alignment coverage.
"""
from __future__ import annotations
import argparse, collections, copy, datetime, gc, gzip, hashlib, io, json, os, resource, time, zipfile
from pathlib import Path
from openpyxl import Workbook, load_workbook
from catalog_json.reconciliation.consensus import build_consensus, case_sha256, receipt_path
from catalog_json.reconciliation.formats import exact_json_equal, resolve_pointer
from catalog_json.reconciliation.materialize import materialize_side
from catalog_json.reconciliation.joint_scope import approve_variant_proposal
from catalog_json.export import to_xlsx
from catalog_json.schema import PageContent

ROLES = ('codex', 'zcode', 'minimax')

def sha(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for block in iter(lambda: f.read(1024 * 1024), b''): h.update(block)
    return h.hexdigest()

def load(path): return json.loads(Path(path).read_text())
def dump(value): return json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + '\n'

def reviewed_cases(task, root):
    """Only explicit manifests; later discussion can replace the SAME full case."""
    chosen = {}
    selection_file = task / 'native-post-source-opinion-selection.json'
    selections = {}
    if selection_file.exists():
        for entry in load(selection_file)['selections']:
            key = (entry['pack_sha256'], entry['role'])
            if key in selections:
                raise ValueError('Duplicate current native opinion selection')
            if entry['role'] not in ROLES or type(entry['round_number']) is not int or entry['round_number'] < 1:
                raise ValueError('Invalid actual native opinion round')
            selections[key] = entry
    for name in ('three-reviewers-ready-manifest.json', 'current-production-ready-manifest.json',
                 'post-source-discussion-ready-manifest.json'):
        path = task / name
        if not path.exists(): continue
        manifest = load(path)
        if tuple(manifest['required_roles']) != ROLES or manifest['standalone_glm']:
            raise ValueError('Only the owner-authorized three-role profile is admissible')
        for item in manifest['packs']:
            pack_path = root / item['path']; pack = load(pack_path)
            if sha(pack_path) != item['file_sha256'] or case_sha256(pack) != item['pack_sha256']:
                raise ValueError('Sealed input pack changed')
            votes = {}
            for role in ROLES:
                number = item['minimax_round'] if role == 'minimax' else item['selected_native_opinion_rounds'][role]
                selected = selections.get((case_sha256(pack), role))
                if selected:
                    if len(pack['cases']) != 1 or selected['case_sha256'] != case_sha256(pack['cases'][0]):
                        raise ValueError('Current opinion must bind the exact unchanged case')
                    number = selected['round_number']
                f = receipt_path(pack, task / 'adjudication', number, role)
                if selected and not f.exists():
                    raise ValueError('Selected current opinion has no saved receipt')
                if not f.exists(): continue
                vote = load(f)
                if selected and selected['receipt_sha256'] != vote['receipt_sha256']:
                    raise ValueError('Selected actual opinion changed')
                # Validates original receipt digest, payload, source/case bindings,
                # identity and actual opinion round, without relabelling it.
                build_consensus(pack, [vote])
                votes[role] = vote
            for case in pack['cases']:
                old = chosen.get(case['case_id'])
                if old and case_sha256(old['case']) != case_sha256(case):
                    raise ValueError('A repeated case ID changed its meaning')
                # An admitted but unfinished discussion cannot erase a completed vote.
                if old and 'minimax' not in votes: continue
                opinions = {role: next(v for v in vote['payload']['verdicts'] if v['case_id'] == case['case_id'])
                            for role, vote in votes.items()}
                signatures = {case_sha256({k: v for k, v in opinion.items() if k in ('decision', 'selected_value')})
                              for opinion in opinions.values()}
                complete = set(votes) == set(ROLES)
                unanimous = complete and len(signatures) == 1
                decision = next(iter(opinions.values()))['decision'] if unanimous else 'unresolved'
                accepted = unanimous and decision in ('ours', 'oleg', 'equivalent', 'both')
                if accepted and any('selected_value' in v for v in opinions.values()):
                    # Replacement JSON requires an explicitly reviewed binding proof;
                    # never ignore a model-proposed value or infer a mapping by search.
                    accepted = False
                chosen[case['case_id']] = {'case': case, 'pack_sha256': case_sha256(pack),
                    'complete': complete, 'accepted': accepted, 'decision': decision,
                    'opinions': opinions, 'votes': votes, 'actual_rounds': {k: v['round_number'] for k, v in votes.items()}}
    for sub in ('armature-111-execution', 'armature-169-row-linkage'):
        folder = root / 'examples/reconciliation' / sub
        pack = load(folder / 'pack.json'); proof = load(folder / 'proof-three-reviewers.json')
        votes = proof['original_votes']
        rebuilt = materialize_side(pack, votes, load(folder / 'bindings.json'))
        if rebuilt != proof: raise ValueError('Published partial proof does not reproduce')
        for case in pack['cases']:
            if case['case_id'] in chosen: raise ValueError('Partial proof duplicates a manifest case')
            result = next(c for c in proof['cases'] if c['case_id'] == case['case_id'])
            chosen[case['case_id']] = {'case': case, 'pack_sha256': case_sha256(pack), 'complete': True,
                'accepted': True, 'decision': result['accepted_side'],
                'opinions': {v['identity']['role']: next(c for c in v['payload']['verdicts'] if c['case_id'] == case['case_id']) for v in votes},
                'votes': {v['identity']['role']: v for v in votes},
                'actual_rounds': {v['identity']['role']: v['round_number'] for v in votes},
                'materialization_proof_sha256': proof['proof_sha256']}
    return list(chosen.values())

def public_record(record):
    case = record['case']
    reviewers = {}
    for role, vote in record['votes'].items():
        ident = vote['identity']
        reviewers[role] = {'role': role, 'session_id': ident['session_id'], 'model': ident['model'],
            'round_number': vote['round_number'], 'receipt_sha256': vote['receipt_sha256'],
            'program_contract_sha256': ident.get('contract_sha256'), 'opinion': record['opinions'][role]}
    return {'case_id': case['case_id'], 'case_sha256': case_sha256(case),
        'pack_sha256': record['pack_sha256'], 'source_sha256': case['source_sha256'],
        'physical_pages': case['physical_pages'], 'accepted': record['accepted'],
        'all_three_present': record['complete'], 'decision': record['decision'],
        'ours': case['ours'], 'oleg': case['oleg'], 'reviewers': reviewers,
        'actual_rounds': record['actual_rounds'],
        'resolved': record.get('resolved', record['accepted']),
        'joint_resolution': record.get('joint_resolution'),
        'materialization_proof_sha256': record.get('materialization_proof_sha256')}

def joint_approvals(task, root, records):
    """Fresh three-party scope proofs resolve parents without rewriting old votes."""
    path = task / 'joint-scope-ready-manifest.json'
    if not path.exists(): return []
    manifest = load(path)
    if tuple(manifest['required_roles']) != ROLES or manifest['standalone_glm']:
        raise ValueError('Unapproved joint review roles')
    parents = {r['case']['case_id']: r for r in records}; result = []; covered = set()
    for item in manifest['packs']:
        file = root / item['path']; pack = load(file)
        if sha(file) != item['file_sha256'] or case_sha256(pack) != item['pack_sha256']:
            raise ValueError('Joint proposal input changed')
        original = [parents[c] for c in item['covered_original_cases']]
        votes = []
        for role in ROLES:
            number = item['minimax_round'] if role == 'minimax' else item['selected_native_opinion_rounds'][role]
            receipt = receipt_path(pack, task / 'adjudication', number, role)
            if receipt.exists():
                vote = load(receipt); build_consensus(pack, [vote]); votes.append(vote)
        if len(votes) != len(ROLES): continue
        consensus = build_consensus(pack, votes)
        if not consensus['all_resolved'] or consensus['cases'][0]['decision'] != 'both': continue
        proof = approve_variant_proposal(pack, votes, original)
        ids = proof['covered_original_case_ids']
        if covered.intersection(ids): raise ValueError('Two joint proposals cover the same parent case')
        covered.update(ids)
        case = pack['cases'][0]
        reviewers = public_record({'case': case, 'pack_sha256': case_sha256(pack),
            'accepted': True, 'complete': True, 'decision': 'both',
            'votes': {v['identity']['role']: v for v in votes},
            'opinions': {v['identity']['role']: v['payload']['verdicts'][0] for v in votes},
            'actual_rounds': {v['identity']['role']: v['round_number'] for v in votes}})['reviewers']
        result.append({'proof_sha256': case_sha256(proof), 'pack_sha256': case_sha256(pack),
            'case_id': case['case_id'], 'source_sha256': case['source_sha256'],
            'physical_pages': case['physical_pages'], 'proposal': proof['proposal'],
            'variant_facts': case['ours'], 'model_option_facts': case['oleg'],
            'reviewers': reviewers, 'catalogue_values_changed': False})
    return result

def apply_joint_resolutions(records, approvals):
    """A scope agreement is explicit additional evidence, never a changed old vote."""
    parents = {r['case']['case_id']: r for r in records}; dispositions = {}
    for approval in approvals:
        for item in approval['proposal']['covered_original_cases']:
            record = parents[item['case_id']]
            if case_sha256(record['case']) != item['case_sha256'] or record['pack_sha256'] != item['pack_sha256']:
                raise ValueError('Scope proof does not bind this retained parent')
            if record.get('joint_resolution'): raise ValueError('Duplicate joint resolution')
            record['resolved'] = True
            record['joint_resolution'] = {'proof_sha256': approval['proof_sha256'],
                'case_id': approval['case_id'], 'decision': 'both',
                'action': 'Keep variant cells unchanged; retain separately bound model options',
                'reviewers': approval['reviewers']}
            for atom in record['case']['ours'] + record['case']['oleg']:
                dispositions.setdefault(atom['fact_id'], []).append({'case_id': item['case_id'],
                    'decision': 'both', 'joint_proof_sha256': approval['proof_sha256']})
    return dispositions

def apply_accepted(bundle, records):
    changes = []; protected = {}; dispositions = {}
    for record in records:
        # The fresh whole-proposal agreement supersedes a parent's earlier
        # scalar choice. Its explicit action is to keep every variant intact.
        if record.get('joint_resolution'): continue
        if not record['accepted']: continue
        case = record['case']; decision = record['decision']
        for atom in case['ours'] + case['oleg']:
            fid = atom.get('fact_id')
            if fid: dispositions.setdefault(fid, []).append({'case_id': case['case_id'], 'decision': decision})
        if decision != 'oleg': continue
        for atom in case['ours']:
            pointer = atom.get('pointer', '')
            if not pointer.startswith('/catalogs/') or '/cells/' not in pointer:
                raise ValueError('No automatic write outside a verified source cell')
            links = case.get('original_ledger_links')
            if links:
                target_ids = next(x['oleg_fact_ids'] for x in links if x['ours_fact_id'] == atom['fact_id'])
                alternatives = [b for b in case['oleg'] if b['fact_id'] in target_ids]
            else:
                # All candidates must give exactly the same value; no fuzzy
                # designation/header matching and no candidate-order winner.
                alternatives = case['oleg']
            if not alternatives or len({case_sha256(b['value']) for b in alternatives}) != 1:
                raise ValueError('Ambiguous replacement; explicit variant binding required')
            new_value = alternatives[0]['value']; cell = resolve_pointer(bundle, pointer)
            expected = atom['value']
            if pointer in protected:
                if not exact_json_equal(protected[pointer], new_value): raise ValueError('Conflicting accepted writes')
                continue
            if not exact_json_equal(cell['value'], expected): raise ValueError('Original target value/type changed')
            if exact_json_equal(expected, new_value): continue
            parts = pointer.split('/'); catalog = bundle['catalogs'][int(parts[2])]
            page = catalog['pages'][int(parts[4])]
            if page['source_sha256'] != case['source_sha256'] or page['pdf_page'] not in case['physical_pages']:
                raise ValueError('Source/page mismatch at target')
            old_cell = copy.deepcopy(cell); cell['value'] = new_value
            state = alternatives[0].get('raw', {}).get('value_state', alternatives[0].get('raw_semantic_context', {}).get('value_state'))
            cell['state'] = 'blank' if state == 'source_blank' and new_value in (None, '') else 'value'
            page.setdefault('reconciliation', {'original_content_sha256': page['content_canonical_sha256'], 'cases': []})['cases'].append(case['case_id'])
            page['reconciliation']['verification_scope'] = 'Three reviewers verified changed fields; earlier full-page review is retained with its actual original hash, not reassigned.'
            page['reviewed_current_content'] = False
            page['content_canonical_sha256'] = case_sha256(page['content'])
            changes.append({'case_id': case['case_id'], 'source_sha256': case['source_sha256'],
                'catalog': catalog['id'], 'physical_page': page['pdf_page'], 'json_pointer': pointer,
                'before': old_cell, 'after': copy.deepcopy(cell), 'selected_input_fact_ids': [b.get('fact_id') for b in alternatives],
                'reviewers': public_record(record)['reviewers']})
            protected[pointer] = copy.deepcopy(new_value)
    return changes, dispositions

def build(root, task, out, *, allow_pending=False):
    start = time.monotonic(); records = reviewed_cases(task, root)
    approvals = joint_approvals(task, root, records)
    joint_dispositions = apply_joint_resolutions(records, approvals)
    alignment = load(task / 'comparison-property-scope-safe/summary.json')
    pending = [r for r in records if not r.get('resolved', r['accepted'])]
    full = alignment.get('semantic_alignment_complete') is True and not pending
    if not full and not allow_pending:
        raise ValueError('Final delivery refused: source-bound alignment and/or three-party adjudication are still incomplete')
    original = root / 'examples/catalogs/all-pages-2026-10-14T172333Z/all-catalogues-980-pages.zip'
    if sha(original) != 'dcfdf4a64448a727766cc50ba31b01cf967b475744b9a6d6c5d065344c5061a7': raise ValueError('Baseline archive SHA changed')
    with zipfile.ZipFile(original) as z: bundle = json.loads(z.read('all-catalogues-980-pages.json'))
    changes, dispositions = apply_accepted(bundle, records)
    for fid, entries in joint_dispositions.items(): dispositions.setdefault(fid, []).extend(entries)
    count = sum(len(c['pages']) for c in bundle['catalogs'])
    if count != 980: raise ValueError('All 980 pages must survive reconciliation')
    for catalog in bundle['catalogs']:
        for page in catalog['pages']: PageContent.model_validate(page['content'])
    # Review coverage and model accuracy are different quantities.
    summary = {'schema_version': 'catalogue-reconciliation-snapshot/1',
        'created_at': datetime.datetime.now(datetime.timezone.utc).isoformat(), 'final_consensus_merge': full,
        'required_reviewers': list(ROLES), 'page_count': count,
        'retained_case_count': len(records), 'three_party_accepted_cases': sum(r['accepted'] for r in records),
        'three_party_resolved_cases': len(records)-len(pending),
        'joint_scope_approved_groups': len(approvals),
        'joint_scope_resolved_cases': sum(bool(r.get('joint_resolution')) for r in records),
        'pending_cases': len(pending), 'canonical_cells_changed': len(changes),
        'alignment_candidates_not_proven_errors': alignment['alignment_counts'],
        'total_real_discrepancies': None, 'semantic_alignment_complete': alignment['semantic_alignment_complete'],
        'direct_disagreement_rate': load(task / 'direct-disagreement-rate.json') if (task / 'direct-disagreement-rate.json').exists() else None,
        'input_ours_zip_sha256': sha(original), 'input_oleg_zip_sha256': sha(task/'input/forenergo-our-extraction.zip'),
        'builder_source_sha256': sha(__file__), 'model_calls_during_export': 0,
        'warning': 'Промежуточное сведение: итоговое количество реальных расхождений ещё не установлено. Непроверенные наблюдения Олега сохранены с происхождением; они не выданы за принятые факты.' if not full else None}
    bundle['baseline_independently_accepted_pages'] = bundle['independently_accepted_pages']
    bundle['independently_accepted_pages'] = sum(bool(v.get('reviewed_current_content')) for c in bundle['catalogs'] for v in c['pages'])
    for catalog in bundle['catalogs']:
        catalog['counts']['independently_accepted'] = sum(bool(v.get('reviewed_current_content')) for v in catalog['pages'])
        catalog['complete_for_selection'] = all(bool(v.get('reviewed_current_content')) for v in catalog['pages'])
        catalog['catalog_complete'] = catalog['complete_for_selection']
    bundle['reconciliation'] = summary; bundle['snapshot_stage'] = 'reconciled_complete' if full else 'reconciliation_partial'
    bundle['complete_for_selection'] = False if not full else bundle['complete_for_selection']
    bundle['reconciliation_cases'] = [public_record(r) for r in records]
    bundle['joint_scope_approvals'] = approvals
    bundle['corrections'] = changes
    out.mkdir(parents=True, exist_ok=False)
    workbook_paths = []
    for catalog in bundle['catalogs']:
        # Existing standard export, including texts, drawings and source producers.
        file = out / (catalog['id']+'.xlsx'); to_xlsx(catalog, file, allow_unverified=True)
        book = load_workbook(file); tab = book.create_sheet('Сведение и исправления')
        tab.append(['Страница PDF','Проверяемый факт','До','После','Codex','ZCode','MiniMax','Статус'])
        for change in changes:
            if change['catalog'] != catalog['id']: continue
            tab.append([change['physical_page'], change['json_pointer'], str(change['before']['value']),
                str(change['after']['value']), *[change['reviewers'][r]['model'] for r in ROLES], 'Три согласованных мнения по изменённому полю'])
        tab.append([None,'Полный консенсус по всем расхождениям',None,None,None,None,None,full])
        options = book.create_sheet('Варианты модели')
        options.append(['Страница PDF','Модель','Характеристика','Исходный список Олега','Значение варианта','JSON указатель варианта','Доказательство','Codex','ZCode','MiniMax'])
        for approval in approvals:
            if not any(p['source_sha256'] == approval['source_sha256'] for p in catalog['pages']): continue
            atoms = {a['fact_id']: a for a in approval['variant_facts'] + approval['model_option_facts']}
            for binding in approval['proposal']['bindings']:
                option = atoms[binding['option_list_fact_id']]
                for member in binding['members']:
                    variant = atoms[member['variant_fact_id']]
                    options.append([approval['physical_pages'][0], binding['model'], binding['property'],
                        option['value'], variant['value'], variant['pointer'], approval['proof_sha256'],
                        *[approval['reviewers'][r]['model'] for r in ROLES]])
        book.save(file); del book; gc.collect(); workbook_paths.append(file)
    archive = out / 'catalogues-unified.zip'; manifest_files = {}
    with zipfile.ZipFile(archive, 'x', compression=zipfile.ZIP_DEFLATED, compresslevel=6) as z:
        for file in workbook_paths:
            z.write(file,file.name); manifest_files[file.name]={'sha256':sha(file),'bytes':file.stat().st_size}
        # A single readable catalogue JSON. Other-side observations use one
        # common literal fact schema and link to the same source/page identities.
        h=hashlib.sha256(); size=0
        with z.open('catalogue.json','w',force_zip64=True) as stream:
            def put(b):
                nonlocal size
                stream.write(b); h.update(b); size+=len(b)
            put(b'{')
            for i,(key,value) in enumerate(bundle.items()):
                if i:put(b',')
                put(json.dumps(key).encode()+b':')
                for piece in json.JSONEncoder(ensure_ascii=False,allow_nan=False,separators=(',',':')).iterencode(value):put(piece.encode())
            put(b',"oleg_observations":['); n=0
            with gzip.open(task/'comparison-property-scope-safe/facts.jsonl.gz','rt') as f:
                for line in f:
                    fact=json.loads(line)
                    if fact['side']!='oleg':continue
                    fact['reconciliation_dispositions']=dispositions.get(fact['fact_id'],[])
                    fact['independently_accepted_in_merged_output']=bool(fact['reconciliation_dispositions']) and all(x['decision'] in ('oleg','equivalent','both') for x in fact['reconciliation_dispositions'])
                    if n:put(b',')
                    put(json.dumps(fact,ensure_ascii=False,separators=(',',':'),allow_nan=False).encode());n+=1
            put(b'],"oleg_original_tables":{')
            for ti,name in enumerate(('sources.json','coverage.json','product_groups.jsonl','products.jsonl','characteristics.jsonl','group_characteristics.jsonl','uncertainties.jsonl')):
                if ti:put(b',')
                put(json.dumps(name).encode()+b':')
                file=task/'input/oleg'/name
                if name.endswith('.jsonl'):
                    put(b'['); first=True
                    with file.open() as original_rows:
                        for line in original_rows:
                            if not line.strip():continue
                            value=json.loads(line)
                            if not first:put(b',')
                            put(json.dumps(value,ensure_ascii=False,separators=(',',':'),allow_nan=False).encode());first=False
                    put(b']')
                else:put(json.dumps(load(file),ensure_ascii=False,separators=(',',':'),allow_nan=False).encode())
            put(b'}}')
        manifest_files['catalogue.json']={'sha256':h.hexdigest(),'bytes':size,'oleg_observations':n}
        reports={'corrections.json':changes, 'reconciliation-report.json':{'summary':summary,'cases':[public_record(r) for r in records], 'joint_scope_approvals': approvals},
            'README.txt':'Единый снимок пяти каталогов (980 физических страниц).\n'+('Итоговый консенсус завершён.\n' if full else 'ВНИМАНИЕ: промежуточный снимок; полный консенсус ещё не завершён.\n')+
            'catalogue.json — основная постраничная выдача, исправления, три мнения по каждому разобранному случаю и сохранённые наблюдения Олега.\nПять Excel содержат те же таблицы, тексты, чертежи, происхождение и исправления.\nВ отчёте отдельно показаны принятые случаи и оставшиеся споры. Кандидаты сопоставления не являются числом ошибок.\n'}
        for name,value in reports.items():
            data=(value if isinstance(value,str) else dump(value)).encode();z.writestr(name,data);manifest_files[name]={'sha256':hashlib.sha256(data).hexdigest(),'bytes':len(data)}
        z.writestr('manifest.json',dump({'schema_version':'catalogue-reconciliation-manifest/1','summary':summary,'files':manifest_files}))
    # Readback needs only the actual archive. Avoid retaining two decoded copies
    # of all980 pages alongside the Oleg tables under the2GiB export budget.
    del bundle
    gc.collect()
    # Read back the actual consumer archive, CRC and every manifest digest.
    with zipfile.ZipFile(archive) as z:
        if z.testzip() is not None:raise ValueError('ZIP CRC failure')
        for name,info in manifest_files.items():
            with z.open(name) as f:
                hh=hashlib.sha256(); observed=0
                for block in iter(lambda:f.read(1024*1024),b''):hh.update(block);observed+=len(block)
            if hh.hexdigest()!=info['sha256'] or observed!=info['bytes']:raise ValueError('Archive manifest mismatch')
        decoded=json.loads(z.read('catalogue.json'))
        if sum(len(c['pages']) for c in decoded['catalogs'])!=980 or len(decoded['oleg_observations'])!=n:raise ValueError('Consumer JSON coverage failed')
        for name,actual in decoded['oleg_original_tables'].items():
            source=task/'input/oleg'/name
            expected=[json.loads(line) for line in source.read_text().splitlines() if line.strip()] if name.endswith('.jsonl') else load(source)
            if not exact_json_equal(actual,expected):raise ValueError('Oleg original table failed typed roundtrip')
            del expected
        checked_cells=0
        for catalog in decoded['catalogs']:
            book=load_workbook(io.BytesIO(z.read(catalog['id']+'.xlsx')),read_only=True,data_only=False)
            expected=[{cell['column_id']:cell['value'] for cell in row['cells']}[column['id']] for page in catalog['pages'] for table in page['content']['tables'] for row in table['rows'] for column in table['columns']]
            actual=[r[4] for r in book['Данные'].iter_rows(min_row=2,values_only=True)]
            if len(expected)!=len(actual) or any(a!=(None if b=='' else b) or (b is not None and b!='' and type(a) is not type(b)) for a,b in zip(actual,expected)):raise ValueError('Excel differs from canonical JSON')
            checked_cells+=len(actual)
            expected_options=[]
            for approval in decoded['joint_scope_approvals']:
                if not any(p['source_sha256']==approval['source_sha256'] for p in catalog['pages']): continue
                atoms={a['fact_id']:a for a in approval['variant_facts']+approval['model_option_facts']}
                for binding in approval['proposal']['bindings']:
                    option=atoms[binding['option_list_fact_id']]
                    for member in binding['members']:
                        variant=atoms[member['variant_fact_id']]
                        expected_options.append((approval['physical_pages'][0],binding['model'],binding['property'],
                            option['value'],variant['value'],variant['pointer'],approval['proof_sha256'],
                            *[approval['reviewers'][r]['model'] for r in ROLES]))
            actual_options=list(book['Варианты модели'].iter_rows(min_row=2,values_only=True))
            if actual_options!=expected_options:raise ValueError('Excel lost model-option bindings or reviewer identities')
            book.close()
    receipt={'archive':str(archive.resolve()),'sha256':sha(archive),'bytes':archive.stat().st_size,
        'page_count':980,'canonical_cells_changed':len(changes),'excel_cells_roundtrip_checked':checked_cells,
        'oleg_observations_retained':n,'archive_crc':'passed','manifest_all_members_verified':True,
        'elapsed_seconds':round(time.monotonic()-start,3),'peak_rss_kib':resource.getrusage(resource.RUSAGE_SELF).ru_maxrss,
        'final_consensus_merge':full,'model_calls':0}
    (out/'verification.json').write_text(dump(receipt));return receipt

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--task',type=Path,default=Path('.tmp/oleg-consensus-20261014'));parser.add_argument('--output',type=Path,required=True);parser.add_argument('--allow-pending',action='store_true');args=parser.parse_args()
    resource.setrlimit(resource.RLIMIT_AS,(2*1024**3,2*1024**3));os.sched_setaffinity(0,{min(os.sched_getaffinity(0))})
    print(dump(build(Path.cwd(),args.task,args.output,allow_pending=args.allow_pending)))
if __name__=='__main__':main()
