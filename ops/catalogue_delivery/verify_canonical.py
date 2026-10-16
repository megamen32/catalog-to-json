#!/usr/bin/env python3
"""Read back the delivered JSON/XLSX; never import the exporter or rewrite data.

This verifies representation, coverage and unresolved/conflicting selections. It
does not certify the accuracy of an extraction against the original PDF.
"""
from __future__ import annotations

import argparse
import collections
import hashlib
import json
from pathlib import Path
import zipfile

CORE = ("id", "catalog", "page", "kind", "model", "parameter", "part",
        "condition", "value", "unit", "state", "source_fact_ids", "producer_ref")
LITERAL = {"id", "catalog", "page", "kind"}
DATA_SHEETS = ("Характеристики", "Чертежи", "Описание")
LONG_MARKER = "[Полное значение: Длинные значения]"
PAGE_COUNTS = {"opn": 60, "isolators": 162, "small": 228,
               "innovations": 216, "armature": 314}


def encoded(value):
    """Type-sensitive comparison (unlike Python's True == 1)."""
    return json.dumps(value, ensure_ascii=False, allow_nan=False, sort_keys=True,
                      separators=(",", ":"))


def digest(path):
    h = hashlib.sha256()
    with open(path, "rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def _json_decode(value):
    def invalid_number(value):
        raise ValueError("Non-finite JSON number: " + value)
    def unique_keys(pairs):
        result = {}
        for key, value in pairs:
            if key in result:
                raise ValueError("Duplicate JSON object key: " + key)
            result[key] = value
        return result
    return json.loads(value, parse_constant=invalid_number, object_pairs_hook=unique_keys)


def _json_load(path):
    with open(path, encoding="utf-8") as stream:
        return _json_decode(stream.read())


def _sheet_for(record):
    if record["kind"] in {"model_parameter", "group_parameter", "model_designation"}:
        return DATA_SHEETS[0]
    return DATA_SHEETS[1] if record["kind"] == "figure" else DATA_SHEETS[2]


def verify(canonical_json, xlsx, *, unresolved_reference=None, manifest=None,
           archive=None, expected_table_ids=None, expected_page_counts=None,
           closure_ledger=None):
    """Return a bounded report; ``ok`` requires every requested proof to pass.

    expected_table_ids is an independently retained JSON list of selected source
    fact IDs (or a dict with source_fact_ids); it must not be derived by this
    verifier from the canonical output being checked.
    """
    failures = []
    failure_count = 0
    def fail(code, **details):
        nonlocal failure_count
        failure_count += 1
        if len(failures) < 60:
            failures.append({"code": code, **details})

    report = {"schema": "canonical-readback-verification/1", "ok": False,
              "scope": "merge selection and JSON/XLSX representation; not PDF accuracy"}
    try:
        doc = _json_load(canonical_json)
        if not isinstance(doc, dict):
            raise ValueError("Canonical document must be an object")
        records = doc["records"]
        if not isinstance(records, list):
            raise ValueError("records must be a list")
    except (OSError, ValueError, KeyError, TypeError) as exc:
        return {**report, "failure_count": 1,
                "failures": [{"code": "json_open", "error": str(exc)}]}
    counts = PAGE_COUNTS if expected_page_counts is None else expected_page_counts
    if doc.get("schema", doc.get("schema_version")) != "canonical-catalogue/1":
        fail("schema")
    sources = doc.get("sources", [])
    source_ids = [s.get("id") for s in sources if isinstance(s, dict)]
    if len(source_ids) != len(set(source_ids)):
        fail("duplicate_sources")
    actual_counts = {s.get("id"): s.get("page_count") for s in sources if isinstance(s, dict)}
    if actual_counts != counts:
        fail("source_page_counts", actual=actual_counts, expected=counts)
    for source in sources:
        if (not isinstance(source, dict) or not isinstance(source.get("filename"), str)
                or not isinstance(source.get("sha256"), str)
                or len(source.get("sha256", "")) != 64
                or any(c not in "0123456789abcdef" for c in source.get("sha256", ""))):
            fail("source_identity")
    wanted_pages = {(catalog, page) for catalog, n in counts.items()
                    for page in range(1, n + 1)}
    pages = doc.get("pages")
    if not isinstance(pages, list):
        fail("missing_pages_manifest")
        page_set = set()
    else:
        page_set = set()
        for page in pages:
            if (not isinstance(page, dict) or type(page.get("page")) is not int
                    or not isinstance(page.get("catalog"), str)):
                fail("invalid_page_manifest_entry")
                continue
            key = (page["catalog"], page["page"])
            if key in page_set:
                fail("duplicate_page_manifest", page=list(key))
            page_set.add(key)
        if page_set != wanted_pages:
            fail("page_manifest_coverage", missing=len(wanted_pages-page_set),
                 extra=len(page_set-wanted_pages))

    producers = doc.get("producers", {})
    if not isinstance(producers, dict):
        fail("producers_object")
        producers = {}
    by_id, physical, source_facts = {}, {}, set()
    catalog_counts, record_pages = collections.Counter(), set()
    for index, record in enumerate(records):
        if not isinstance(record, dict) or any(k not in record for k in CORE):
            fail("record_fields", index=index)
            continue
        if set(record) != set(CORE):
            fail("unrepresented_record_fields", index=index,
                 fields=sorted(set(record)-set(CORE)))
        ident = record["id"]
        if not isinstance(ident, str) or not ident:
            fail("record_id", index=index)
            continue
        if ident in by_id:
            fail("duplicate_record_id", id=ident)
        by_id[ident] = record
        if any(not isinstance(record[k], str) for k in ("catalog", "kind")):
            fail("record_identity_type", id=ident)
            continue
        catalog_counts[record["catalog"]] += 1
        if type(record["page"]) is not int or (record["catalog"], record["page"]) not in wanted_pages:
            fail("record_source_page", id=ident)
        else:
            record_pages.add((record["catalog"], record["page"]))
        refs = record["source_fact_ids"]
        if not isinstance(refs, list) or not refs or any(not isinstance(r, str) or not r for r in refs):
            fail("source_fact_ids", id=ident)
        elif len(refs) != len(set(refs)):
            fail("duplicate_source_fact_id_in_record", id=ident)
        else:
            source_facts.update(refs)
        producer = record["producer_ref"]
        if not isinstance(producer, str) or producer not in producers or not producers.get(producer):
            fail("missing_producer", id=ident)
        if record["kind"] in {"model_parameter", "group_parameter"}:
            # Page is evidence, never a fabricated applicability condition.
            key = encoded([record[k] for k in ("catalog", "kind", "model", "parameter", "part", "condition")])
            value = encoded([record[k] for k in ("value", "unit", "state")])
            if key in physical:
                other_value, other_id = physical[key]
                fail("physical_value_conflict" if other_value != value else "duplicate_physical_characteristic",
                     ids=[other_id, ident])
            else:
                physical[key] = value, ident

    original_unresolved = None
    if unresolved_reference is None:
        fail("unresolved_reference_required")
    else:
        try:
            original = _json_load(unresolved_reference)
            original_unresolved = original["facts"] if isinstance(original, dict) else original
            if len(original_unresolved) != 42:
                fail("reference_unresolved_count", count=len(original_unresolved))
            if closure_ledger is None:
                if doc.get("metadata", {}).get("source_review"):
                    fail("closure_ledger_required")
                if encoded(doc.get("unresolved")) != encoded(original_unresolved):
                    fail("retained_unresolved_mismatch")
            else:
                ledger = _json_load(closure_ledger)
                review = doc.get("metadata", {}).get("source_review", {})
                if ledger.get("schema") != "canonical-source-review-ledger/1":
                    fail("closure_ledger_schema")
                sha = hashlib.sha256(encoded(ledger).encode()).hexdigest()
                if sha != review.get("ledger_content_sha256"):
                    fail("closure_ledger_identity")
                if ledger.get("input_sha256") != review.get("input_sha256"):
                    fail("closure_input_identity")
                if (hashlib.sha256(encoded(records).encode()).hexdigest() != ledger.get("output_records_sha256")
                        or len(records) != ledger.get("output_record_count")):
                    fail("reviewed_recordset_mismatch")
                if encoded(ledger.get("original_unresolved")) != encoded(original_unresolved):
                    fail("closure_original_history_changed")
                old = {u["fact"]["fact_id"]: u for u in original_unresolved}
                decisions = ledger.get("closures", [])
                decision_ids = [c["fact_id"] for c in decisions]
                if len(set(decision_ids)) != len(decision_ids) or set(decision_ids) != set(old):
                    fail("closure_partition")
                closed = set()
                for c in decisions:
                    if c.get("disposition") not in {"closed", "retained"}:
                        fail("closure_disposition")
                    if not c.get("reason") or not c.get("evidence"):
                        fail("closure_evidence")
                    u = old.get(c["fact_id"])
                    if u is None:
                        continue
                    if hashlib.sha256(encoded(u).encode()).hexdigest() != c.get("before_sha256"):
                        fail("closure_before_identity")
                    if c["disposition"] == "closed":
                        closed.add(c["fact_id"])
                    for ident in c.get("record_ids", []):
                        if ident not in by_id:
                            fail("closure_result_missing", id=ident)
                remaining = [u for u in original_unresolved if u["fact"]["fact_id"] not in closed]
                if encoded(doc.get("unresolved")) != encoded(remaining):
                    fail("closure_remaining_mismatch")
                if review.get("closed_count") != len(closed) or review.get("retained_count") != len(remaining):
                    fail("closure_counts")
                for change in ledger.get("record_changes", []):
                    after = change.get("after")
                    if after is not None:
                        if encoded(by_id.get(after["id"])) != encoded(after):
                            fail("reviewed_record_mismatch", id=after["id"])
                    elif change.get("before", {}).get("id") in by_id:
                        fail("reviewed_removal_still_present")
                for change in ledger.get("metadata_changes", []):
                    if encoded(doc.get("metadata", {}).get(change["key"])) != encoded(change["after"]):
                        fail("reviewed_metadata_mismatch", key=change["key"])
            if isinstance(original, dict) and original.get("source_archive_sha256"):
                if doc.get("metadata", {}).get("input_archive_sha256") != original["source_archive_sha256"]:
                    fail("input_archive_identity")
        except (OSError, ValueError, KeyError, TypeError) as exc:
            fail("unresolved_reference_open", error=str(exc))
    if expected_table_ids is None:
        fail("selected_table_reference_required")
    else:
        try:
            table_ids = _json_load(expected_table_ids) if isinstance(expected_table_ids, (str, Path)) else expected_table_ids
            if isinstance(table_ids, dict):
                table_ids = table_ids["source_fact_ids"]
            if not isinstance(table_ids, list) or any(not isinstance(x, str) for x in table_ids):
                raise ValueError("Expected selected-table IDs must be a string list")
            missing = set(table_ids) - source_facts
            if missing:
                fail("selected_table_facts_missing", count=len(missing), sample=sorted(missing)[:5])
            if encoded(doc.get("metadata", {}).get("selected_table_source_fact_ids")) != encoded(table_ids):
                fail("selected_table_metadata_mismatch")
        except (OSError, ValueError, KeyError, TypeError) as exc:
            fail("selected_table_reference", error=str(exc))

    open_counts = collections.Counter(r.get("state") for r in records
                                     if r.get("state") in {"source_conflict", "unknown"})
    xlsx_seen, xlsx_unresolved, xlsx_metadata = set(), [], {}
    current_issue_sheet = False
    try:
        from openpyxl import load_workbook
        book = load_workbook(xlsx, read_only=True, data_only=False)
        chunks = {}
        if "Длинные значения" in book.sheetnames:
            iterator = book["Длинные значения"].iter_rows()
            header = [c.value for c in next(iterator)]
            if header != ["sheet", "row", "record_id", "field", "chunk_index", "chunk_count", "chunk"]:
                raise ValueError("Unexpected long-value sheet headers")
            for row in iterator:
                values = [c.value for c in row]
                if not any(v is not None for v in values):
                    continue
                sheet, number, rid, field, part, total, chunk = values
                if (type(number) is not int or type(part) is not int or type(total) is not int
                        or not 0 <= part < total or not isinstance(chunk, str)):
                    raise ValueError("Invalid long-value chunk")
                if any(c.data_type == "f" for c in row):
                    raise ValueError("Formula in long-value chunk")
                key = sheet, number, field
                info = chunks.setdefault(key, {"rid": rid, "total": total, "parts": {}, "used": False})
                if info["rid"] != rid or info["total"] != total or part in info["parts"]:
                    raise ValueError("Conflicting/duplicate long-value chunks")
                info["parts"][part] = chunk
        def cell_text(cell, sheet, number, field, rid):
            if cell.data_type == "f":
                raise ValueError("Formula in authoritative cell")
            key = sheet, number, field
            if key not in chunks:
                return cell.value
            info = chunks[key]
            if cell.value != LONG_MARKER or info["rid"] != rid:
                raise ValueError("Long-value marker/record ID mismatch")
            if set(info["parts"]) != set(range(info["total"])):
                raise ValueError("Missing long-value chunk")
            info["used"] = True
            return "".join(info["parts"][i] for i in range(info["total"]))
        for sheet in DATA_SHEETS:
            if sheet not in book.sheetnames:
                fail("missing_data_sheet", sheet=sheet)
                continue
            iterator = book[sheet].iter_rows()
            header = [c.value for c in next(iterator)]
            if header != list(CORE) + ["value_display"]:
                raise ValueError("Unexpected data headers: " + sheet)
            for number, row in enumerate(iterator, 2):
                if not any(c.value is not None for c in row):
                    continue
                if any(c.data_type == "f" for c in row):
                    raise ValueError("Formula in data row (including display)")
                ident = row[0].value
                reconstructed = {}
                for field, cell in zip(CORE, row):
                    value = cell_text(cell, sheet, number, field, ident)
                    if field not in LITERAL:
                        if not isinstance(value, str):
                            raise ValueError("Non-JSON typed field: " + field)
                        value = _json_decode(value)
                    elif field == "page":
                        if type(value) is not int:
                            raise ValueError("Excel page must be integer")
                    elif not isinstance(value, str):
                        raise ValueError("Excel identity must be string")
                    reconstructed[field] = value
                if ident in xlsx_seen:
                    fail("xlsx_duplicate_record", id=ident)
                xlsx_seen.add(ident)
                expected = by_id.get(ident)
                if expected is None:
                    fail("xlsx_extra_record", id=ident)
                else:
                    if _sheet_for(expected) != sheet:
                        fail("xlsx_wrong_kind_sheet", id=ident, sheet=sheet)
                    for field in CORE:
                        if encoded(reconstructed[field]) != encoded(expected[field]):
                            fail("xlsx_field_mismatch", id=ident, field=field)
        issue_fields = ("catalog", "page", "model", "parameter", "state", "id")
        issue_headers = ["Каталог", "Страница", "Модель", "Параметр", "state", "id"]
        catalog_names = {"armature": "Арматура", "small": "Малый каталог",
                         "innovations": "Инновационные изделия 2021-2022",
                         "isolators": "Изоляторы", "opn": "ОПН"}
        actual_issues = []
        for sheet, headers in (("Неоднозначные", ["index", "value"]), ("Метаданные", ["key", "value"])):
            if sheet not in book.sheetnames:
                fail("missing_envelope_sheet", sheet=sheet)
                continue
            iterator = book[sheet].iter_rows()
            actual_headers = [c.value for c in next(iterator)]
            if sheet == "Неоднозначные" and actual_headers == issue_headers:
                current_issue_sheet = True
                for number, row in enumerate(iterator, 2):
                    if not any(c.value is not None for c in row):
                        continue
                    ident = row[5].value
                    actual_issues.append([cell_text(cell, sheet, number, field, ident)
                                          for field, cell in zip(issue_fields, row)])
                continue
            if actual_headers != headers:
                raise ValueError("Unexpected envelope headers: " + sheet)
            for number, row in enumerate(iterator, 2):
                if not any(c.value is not None for c in row):
                    continue
                key = row[0].value
                if row[0].data_type == "f":
                    raise ValueError("Formula in envelope key")
                continuation_id = "unresolved:" + str(key) if sheet == "Неоднозначные" else key
                raw = cell_text(row[1], sheet, number, "value", continuation_id)
                value = _json_decode(raw)
                if sheet == "Неоднозначные":
                    if type(key) is not int or key != len(xlsx_unresolved):
                        fail("xlsx_unresolved_index", row=number)
                    xlsx_unresolved.append(value)
                else:
                    if key in xlsx_metadata:
                        fail("xlsx_duplicate_metadata", key=key)
                    xlsx_metadata[key] = value
        for key, info in chunks.items():
            if not info["used"]:
                fail("xlsx_unused_long_value", sheet=key[0], row=key[1], field=key[2])
        book.close()
        if xlsx_seen != set(by_id):
            fail("xlsx_record_coverage", missing=len(set(by_id)-xlsx_seen), extra=len(xlsx_seen-set(by_id)))
        if current_issue_sheet:
            expected_issues = []
            for record in records:
                if record["state"] not in {"source_conflict", "unknown"}:
                    continue
                values = []
                for field in issue_fields:
                    value = record[field]
                    if field == "catalog":
                        value = catalog_names.get(value, value)
                    elif value is not None and not isinstance(value, (str, int, float)):
                        value = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
                    values.append(value)
                expected_issues.append(values)
            if encoded(actual_issues) != encoded(expected_issues):
                fail("xlsx_open_characteristics_mismatch", actual=len(actual_issues), expected=len(expected_issues))
            envelope = {k: v for k, v in doc.items() if k != "records"}
        else:
            if encoded(xlsx_unresolved) != encoded(doc.get("unresolved")):
                fail("xlsx_unresolved_mismatch")
            envelope = {k: v for k, v in doc.items() if k not in {"records", "unresolved"}}
        if encoded(xlsx_metadata) != encoded(envelope):
            fail("xlsx_metadata_mismatch", actual_keys=sorted(xlsx_metadata), expected_keys=sorted(envelope))
    except (OSError, ValueError, KeyError, TypeError, StopIteration, zipfile.BadZipFile) as exc:
        fail("xlsx_open_or_decode", error=str(exc))

    paths = {"catalogue.json": Path(canonical_json), "catalogue.xlsx": Path(xlsx)}
    if manifest is not None:
        try:
            manifest_doc = _json_load(manifest)
            if current_issue_sheet:
                wanted = {"source_conflict_count": open_counts["source_conflict"],
                          "unknown_count": open_counts["unknown"],
                          "unresolved_count": sum(open_counts.values()),
                          "legacy_unresolved_count": len(doc.get("unresolved", []))}
                for key, expected in wanted.items():
                    if manifest_doc.get(key) != expected:
                        fail("manifest_open_count", field=key, expected=expected)
                readme = (Path(manifest).parent / "README.txt").read_text(encoding="utf-8")
                for expected in (
                    f"Открытых противоречий (source_conflict): {open_counts['source_conflict']}.",
                    f"Неизвестных значений (unknown): {open_counts['unknown']}.",
                    f"Всего нерешённых характеристик на листе «Неоднозначные»: {sum(open_counts.values())}.",
                ):
                    if expected not in readme:
                        fail("readme_open_count", expected=expected)
            entries = manifest_doc["files"]
            if set(entries) != {"catalogue.json", "catalogue.xlsx", "README.txt"}:
                fail("manifest_file_set")
            for name, info in entries.items():
                path = paths.get(name, Path(manifest).parent / name)
                paths[name] = path
                if info.get("sha256") != digest(path) or info.get("bytes") != path.stat().st_size:
                    fail("manifest_hash_or_size", file=name)
            paths["manifest.json"] = Path(manifest)
        except (OSError, ValueError, KeyError, TypeError) as exc:
            fail("manifest_open", error=str(exc))
    if archive is not None:
        try:
            if manifest is None:
                fail("archive_manifest_required")
            with zipfile.ZipFile(archive) as zipped:
                names = zipped.namelist()
                if len(names) != 4 or set(names) != {"catalogue.json", "catalogue.xlsx", "README.txt", "manifest.json"}:
                    fail("archive_file_set", files=names[:10])
                if zipped.testzip() is not None:
                    fail("archive_crc")
                for name, path in paths.items():
                    if name not in names:
                        continue
                    h = hashlib.sha256()
                    with zipped.open(name) as stream:
                        for block in iter(lambda: stream.read(1024*1024), b""):
                            h.update(block)
                    if h.hexdigest() != digest(path):
                        fail("archive_file_bytes", file=name)
        except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
            fail("archive_open", error=str(exc))
    return {**report, "ok": failure_count == 0, "failure_count": failure_count,
            "failures": failures, "failures_truncated": failure_count > len(failures),
            "coverage": {"sources": len(source_ids), "pages": len(page_set),
                         "pages_with_records": len(record_pages), "records": len(records),
                         "records_by_catalog": dict(catalog_counts), "xlsx_records": len(xlsx_seen),
                         "unresolved": len(doc.get("unresolved", [])),
                         "source_conflict": open_counts["source_conflict"],
                         "unknown": open_counts["unknown"],
                         "open_characteristics": sum(open_counts.values()),
                         "source_fact_ids": len(source_facts)},
            "sha256": {"catalogue.json": digest(canonical_json),
                       "catalogue.xlsx": digest(xlsx) if Path(xlsx).is_file() else None}}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("canonical_json", type=Path)
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("--unresolved-reference", type=Path, required=True)
    parser.add_argument("--expected-table-ids", type=Path, required=True)
    parser.add_argument("--manifest", type=Path)
    parser.add_argument("--archive", type=Path)
    parser.add_argument("--closure-ledger", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    report = verify(args.canonical_json, args.xlsx, unresolved_reference=args.unresolved_reference,
                    expected_table_ids=args.expected_table_ids, manifest=args.manifest,
                    archive=args.archive, closure_ledger=args.closure_ledger)
    result = json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n"
    if args.output:
        args.output.write_text(result, encoding="utf-8")
    print(result, end="")
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
