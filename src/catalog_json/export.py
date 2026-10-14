import html
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from .schema import PageContent


def accepted_pages(catalog,allow_unverified=False):
    if not allow_unverified and not catalog.get('complete_for_selection'):
        raise ValueError("Catalog selection is not fully reviewed; export refused")
    return [p for p in catalog['pages'] if p.get('content') and (p['status']=='accepted' or allow_unverified)]


def column_heading(c):
    label=(c.group+' / ' if c.group else '')+c.label
    if c.unit and c.unit not in c.label:label+=', '+c.unit
    if c.qualifier:label+='; '+c.qualifier
    return label


def to_xlsx(catalog,path: Path,allow_unverified=False):
    pages=accepted_pages(catalog,allow_unverified);w=Workbook();s=w.active;s.title='Данные'
    s.append(['PDF страница','Таблица','Строка','Столбец','Значение','Единица','Состояние','Группа','Условие'])
    texts=w.create_sheet('Текст и чертежи');texts.append(['PDF страница','Тип','Заголовок','Текст','Подписи'])
    provenance=w.create_sheet('Происхождение')
    provenance.append(['PDF страница','Напечатанный номер','Статус','Исходный PDF','SHA исходного PDF',
                       'Извлекатель','Версия извлекателя','SHA кода извлекателя','Проверяющий','Версия проверяющего','SHA кода проверяющего',
                       'Последнее извлечение, с','Последняя проверка, с'])
    def producer(page,key):
        origin=page.get(key) or {}
        if origin.get('status') not in {'receipt_verified','manual_source_review_verified'}:return ['не подтверждён','не подтверждена',None,None]
        program=origin.get('program',{});elapsed=None
        if all(isinstance(origin.get(k),(int,float)) for k in ('started_at','finished_at')):
            elapsed=round(origin['finished_at']-origin['started_at'],3)
        return [origin.get('model'),program.get('package_version','не подтверждена'),program.get('source_code_sha256'),elapsed]
    for page in pages:
        data=PageContent.model_validate(page['content'])
        extraction=producer(page,'candidate_producer');review=producer(page,'review_producer')
        provenance.append([page['pdf_page'],data.printed_page,page['status'],catalog.get('source',{}).get('filename'),
                           page.get('source_sha256',catalog.get('source',{}).get('sha256')),*extraction[:3],*review[:3],extraction[3],review[3]])
        texts.append([page['pdf_page'],'Заголовок',data.title,'',''])
        for block in data.text_blocks:texts.append([page['pdf_page'],block.kind,'',block.text,''])
        for figure in data.figures:texts.append([page['pdf_page'],'Чертёж',figure.title,figure.description,', '.join(figure.labels)])
        for note in data.source_notes+data.uncertainties:texts.append([page['pdf_page'],'Примечание','',note,''])
        for table_index,table in enumerate(data.tables,1):
            wide=w.create_sheet(f'Стр {page["pdf_page"]} таб {table_index}'[:31]);wide.append([column_heading(c) for c in table.columns])
            for row in table.rows:
                cells={c.column_id:c for c in row.cells};wide.append([cells[c.id].value for c in table.columns])
            wide.freeze_panes='B2';wide.auto_filter.ref=wide.dimensions
            for column in wide.columns:wide.column_dimensions[column[0].column_letter].width=25
            cols={c.id:c for c in table.columns}
            for n,row in enumerate(table.rows,1):
                for cell in row.cells:
                    col=cols[cell.column_id];s.append([page['pdf_page'],table.title,n,col.label,cell.value,col.unit,cell.state,col.group,col.qualifier])
    s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
    for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='244C66')
    for col,width in {'A':16,'B':35,'C':12,'D':65,'E':28,'F':16,'G':18}.items():s.column_dimensions[col].width=width
    for row in s:
        for c in row:
            c.alignment=Alignment(wrap_text=True,vertical='top')
            if isinstance(c.value,str) and c.value.startswith('='):c.data_type='s'
    for sheet in w:
        for c in sheet[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='244C66')
        for row in sheet:
            for c in row:
                c.alignment=Alignment(wrap_text=True,vertical='top')
                if isinstance(c.value,str) and c.value.startswith('='):c.data_type='s'
    for col,width in {'A':16,'B':25,'C':35,'D':100,'E':60}.items():texts.column_dimensions[col].width=width
    provenance.freeze_panes='C2';provenance.auto_filter.ref=provenance.dimensions
    for column in provenance.columns:provenance.column_dimensions[column[0].column_letter].width=30
    w.save(path)


def to_html(catalog,path: Path,allow_unverified=False):
    pages=accepted_pages(catalog,allow_unverified);h=html.escape
    parts=['<!doctype html><meta charset="utf-8"><title>Catalog JSON</title><style>body{font:17px system-ui;max-width:1450px;margin:24px auto;color:#253a49;padding:16px}table{border-collapse:collapse;width:100%;margin-bottom:24px}td,th{padding:10px;border:1px solid #ccc;text-align:left;vertical-align:top}th{background:#e7eff5}details{margin:18px 0}pre{white-space:pre-wrap} .note{background:#fff2cb;padding:15px}</style>']
    for page in pages:
        data=PageContent.model_validate(page['content']);parts.append(f'<h1>{h(data.title)}</h1><p>PDF: {page["pdf_page"]} · статус: {h(page["status"])}</p>')
        for block in data.text_blocks:parts.append(f'<p>{h(block.text)}</p>')
        for table in data.tables:
            parts.append('<h2>'+h(table.title)+'</h2><table><tr>'+''.join('<th>'+h(column_heading(c))+'</th>' for c in table.columns)+'</tr>')
            for row in table.rows:
                cells={c.column_id:c for c in row.cells}
                parts.append('<tr>'+''.join('<td>'+h('' if cells[c.id].value is None else str(cells[c.id].value))+'</td>' for c in table.columns)+'</tr>')
            parts.append('</table>')
        for figure in data.figures:parts.append('<h2>'+h(figure.title)+'</h2><p>'+h(figure.description)+'</p><p>'+h(', '.join(figure.labels))+'</p>')
        for note in data.source_notes+data.uncertainties:parts.append('<p class="note">'+h(note)+'</p>')
    path.write_text(''.join(parts),encoding='utf-8')
